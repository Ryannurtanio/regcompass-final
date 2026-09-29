"""Generate config/indicators.json + config/pillars.json for all 12 RDTII 2.1
pillars from the organizer's own two sheets. LOCAL DEV ONLY (like
make_gt_matrix.py): it needs openpyxl and the organizer workbooks, which do not
ship with the product. Run it with:

    uv run --with openpyxl python scripts/make_indicators.py

Sources and who wins:

* "Indicator Reference" (OUTPUT_TEMPLATE_FINAL_ROUND.xlsx) is the authority for
  the ID STRING, the indicator name, the pillar name, the weight and the
  exception note. Every ID in that sheet is stored as text, so it round-trips.
* "RDTII 2.1 Methodology" (ESCAP-RDTII-2.1_ Round 2 Database.xlsx) is the
  authority for the scoring criteria text and the score levels. 54 of its 61 ID
  cells are FLOATS, so the join is on the normalised numeric value and the
  generator never calls str() on an ID cell.

ID spelling: the same indicator is `4.1` in the Indicator Reference and renders
as `4.10` in the Methodology sheet (a float carrying number format 0.00). We
emit `4.1`, the Indicator Reference spelling, because that is the sheet shipped
with the final-round template and the one the judges read. `4.01`, `12.01` and
`12.4.1` to `12.4.7` are likewise emitted exactly as that sheet prints them.

62 indicators are generated. 6.5 (treaty participation) exists only in the
Indicator Reference and is not answerable from legislation, so it carries
legislation_mapped: false and the Gate skips it.

Definitions: the 9 Round 1 indicators keep their hand-written name and
definition byte for byte (the golden evidence files pin both). The other 53 get
a definition composed mechanically from the sheets, marked
definition_derived: true. No model call of any kind happens here.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PAPERWORK = ROOT.parent

DEFAULT_METHODOLOGY = (
    PAPERWORK / "Knowledge Base" / "Databases" / "ESCAP-RDTII-2.1_ Round 2 Database.xlsx"
)
DEFAULT_REFERENCE = (
    PAPERWORK
    / "Round 3"
    / "original_materials"
    / "Finalist Orientation"
    / "OUTPUT_TEMPLATE_FINAL_ROUND.xlsx"
)

METHODOLOGY_SHEET = "RDTII 2.1 Methodology"
REFERENCE_SHEET = "Indicator Reference"

# Indicators that are not answerable from an economy's legislation: 6.5 is
# participation in an agreement with binding commitments on data transfer,
# which is treaty membership, not a provision anyone can quote from a statute.
NOT_LEGISLATION_MAPPED = frozenset({"6.5"})

# What a Run covers when the reviewer names no Pillar. The organizer's workbook
# header calls 6 and 7 the mandatory pair; all 12 are in scope on request.
DEFAULT_PILLARS = [6, 7]

ID_SPELLING_NOTE = (
    "Indicator IDs are TEXT, spelled as the Indicator Reference sheet spells them:"
    " 4.1 (not 4.10, which is how the Methodology sheet's float renders under its"
    " 0.00 number format), 4.01, 12.01 and 12.4.1 to 12.4.7."
)


# ---------------------------------------------------------------------------
# reading the sheets
# ---------------------------------------------------------------------------


def _norm_key(value) -> str:
    """The join key between the two sheets. Three-level IDs are text in both
    sheets and join as text; every other ID is a float in the Methodology sheet
    and text in the Indicator Reference, so both sides normalise through float.
    str() is never applied to a Methodology ID cell."""
    if isinstance(value, str) and value.count(".") == 2:
        return value.strip()
    return f"{round(float(value), 6):.6f}"


def _text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _weight(value) -> float | None:
    """The Indicator Reference stores weights as the text '38%'."""
    if value is None:
        return None
    if isinstance(value, (int, float)):
        return round(float(value), 6)
    text = str(value).strip()
    if not text:
        return None
    if text.endswith("%"):
        return round(float(text[:-1].strip()) / 100.0, 6)
    return round(float(text), 6)


def read_reference(path: Path) -> tuple[list[dict], dict[int, str]]:
    """Rows of the Indicator Reference sheet, plus pillar number -> pillar name
    read off the banner rows ('Pillar 4: Intellectual Property Rights')."""
    import openpyxl

    ws = openpyxl.load_workbook(path, data_only=True)[REFERENCE_SHEET]
    rows: list[dict] = []
    pillar_names: dict[int, str] = {}
    pillar = None
    for r in range(4, ws.max_row + 1):
        banner = _text(ws.cell(r, 1).value)
        if banner and banner.lower().startswith("pillar "):
            head, _, name = banner.partition(":")
            pillar = int(head.split()[1])
            pillar_names[pillar] = name.strip()
            continue
        indicator_id = _text(ws.cell(r, 2).value)
        name = _text(ws.cell(r, 3).value)
        if not indicator_id or not name:
            continue  # the trailing mapping-traps block carries no ID
        if pillar is None:
            raise ValueError(f"{REFERENCE_SHEET} row {r}: indicator before any pillar banner")
        rows.append(
            {
                "id": indicator_id,
                "pillar": pillar,
                "name": name,
                "weight": _weight(ws.cell(r, 4).value),
                "exception": _text(ws.cell(r, 5).value),
                "row": r,
            }
        )
    return rows, pillar_names


def read_methodology(path: Path) -> dict[str, dict]:
    """Join key -> {criteria, score_levels} from the Methodology sheet."""
    import openpyxl

    ws = openpyxl.load_workbook(path, data_only=True)[METHODOLOGY_SHEET]
    out: dict[str, dict] = {}
    for r in range(2, ws.max_row + 1):
        raw_id = ws.cell(r, 2).value
        criteria = _text(ws.cell(r, 4).value)
        if raw_id is None or criteria is None:
            continue  # pillar banner rows carry the pillar name in column B
        levels = _text(ws.cell(r, 5).value) or ""
        out[_norm_key(raw_id)] = {
            "criteria": criteria,
            "score_levels": [ln.strip() for ln in levels.splitlines() if ln.strip()],
            "row": r,
        }
    return out


# ---------------------------------------------------------------------------
# composing the registry
# ---------------------------------------------------------------------------


def _derived_definition(name: str, criteria: str | None, exception: str | None) -> str:
    parts = [name.rstrip(".") + "."]
    if criteria:
        parts.append("Scoring criteria: " + " ".join(criteria.split()))
    if exception:
        parts.append(" ".join(exception.split()))
    return " ".join(parts)


def build(methodology_path: Path, reference_path: Path, existing: dict) -> tuple[dict, dict]:
    rows, pillar_names = read_reference(reference_path)
    methodology = read_methodology(methodology_path)

    indicators: dict[str, dict] = {}
    by_pillar: dict[int, list[str]] = {}
    for row in rows:
        indicator_id = row["id"]
        if int(indicator_id.split(".")[0]) != row["pillar"]:
            raise ValueError(
                f"indicator {indicator_id} sits under pillar banner {row['pillar']}"
            )
        method = methodology.get(_norm_key(indicator_id), {})
        criteria = method.get("criteria")
        exception = row["exception"]
        kept = existing.get(indicator_id)
        if kept is not None:
            name = kept["name"]
            definition = kept["definition"]
            derived = False
        else:
            name = row["name"]
            definition = _derived_definition(name, criteria, exception)
            derived = True
        indicators[indicator_id] = {
            "pillar": row["pillar"],
            "name": name,
            "definition": definition,
            "definition_derived": derived,
            "criteria": criteria,
            "score_levels": method.get("score_levels", []),
            "weight": row["weight"],
            "exception": exception,
            "legislation_mapped": indicator_id not in NOT_LEGISLATION_MAPPED,
        }
        by_pillar.setdefault(row["pillar"], []).append(indicator_id)

    pillars = {
        str(p): {"name": pillar_names[p], "indicator_ids": by_pillar[p]}
        for p in sorted(by_pillar)
    }
    return indicators, pillars


def fixture_extract(indicators: dict, methodology_path: Path) -> dict:
    """The committed organizer extract the acceptance test reads, so the test
    never needs the xlsx: the id set and count of BOTH sheets."""
    methodology = read_methodology(methodology_path)
    return {
        "_comment": (
            "Extract of the organizer's two indicator sheets, written by"
            " scripts/make_indicators.py. reference_ids is the Indicator Reference"
            " sheet (62 rows, the authority for the ID spelling); methodology_ids"
            " is the RDTII 2.1 Methodology sheet (61 rows) expressed in the same"
            " spelling through the numeric join."
        ),
        "reference_ids": list(indicators),
        "reference_count": len(indicators),
        "methodology_ids": [
            i for i in indicators if _norm_key(i) in methodology
        ],
        "methodology_count": len(methodology),
        "pillars": {
            str(p): [i for i, e in indicators.items() if e["pillar"] == p]
            for p in sorted({e["pillar"] for e in indicators.values()})
        },
    }


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--methodology", type=Path, default=DEFAULT_METHODOLOGY)
    ap.add_argument("--reference", type=Path, default=DEFAULT_REFERENCE)
    ap.add_argument("--out", type=Path, default=ROOT / "config" / "indicators.json")
    ap.add_argument("--pillars-out", type=Path, default=ROOT / "config" / "pillars.json")
    ap.add_argument(
        "--fixture-out",
        type=Path,
        default=ROOT / "tests" / "fixtures" / "organizer_indicators.json",
    )
    args = ap.parse_args()

    for path in (args.methodology, args.reference):
        if not path.is_file():
            raise SystemExit(f"organizer workbook not found: {path}")

    # Only HAND-WRITTEN entries are carried over. An entry this script derived on
    # an earlier run is regenerated, so running the generator twice is a no-op
    # rather than a promotion of derived text into hand-written text.
    existing_raw = json.loads(args.out.read_text(encoding="utf-8")) if args.out.exists() else {}
    existing = {
        k: v
        for k, v in existing_raw.items()
        if not k.startswith("_") and not v.get("definition_derived", False)
    }

    indicators, pillars = build(args.methodology, args.reference, existing)

    payload = {
        "_generated_by": "scripts/make_indicators.py",
        "_sources": [
            f"{args.methodology.name}#'{METHODOLOGY_SHEET}' (criteria, score levels)",
            f"{args.reference.name}#'{REFERENCE_SHEET}' (id, name, pillar, weight, exception)",
        ],
        "_id_spelling_note": ID_SPELLING_NOTE,
        "_definition_note": (
            "definition_derived: false means the definition and the name are the"
            " hand-written Round 1 text, kept byte for byte. true means both were"
            " composed mechanically from the two sheets, never by a model."
        ),
        **indicators,
    }
    args.out.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    pillars_payload = {
        "_generated_by": "scripts/make_indicators.py",
        "_comment": (
            "Pillar number -> official pillar name and the indicator IDs under it,"
            " read off the Indicator Reference banner rows. The pillar DESCRIPTION"
            " the Gate embeds lives in config/pillar_N_keywords.json."
        ),
        "_default_pillars_note": (
            "default_pillars is what a Run covers when the reviewer names no"
            " Pillar. Pillars 6 and 7 are the mandatory pair (the organizer's"
            " workbook header), and defaulting to all 12 would turn every"
            " unqualified Run into a 61-Indicator sweep nobody asked for."
        ),
        "default_pillars": DEFAULT_PILLARS,
        **pillars,
    }
    args.pillars_out.write_text(
        json.dumps(pillars_payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )

    args.fixture_out.write_text(
        json.dumps(fixture_extract(indicators, args.methodology), indent=2, ensure_ascii=False)
        + "\n",
        encoding="utf-8",
    )

    print(f"{len(indicators)} indicators across {len(pillars)} pillars -> {args.out}")
    print(f"pillars -> {args.pillars_out}")
    print(f"organizer extract -> {args.fixture_out}")


if __name__ == "__main__":
    main()
