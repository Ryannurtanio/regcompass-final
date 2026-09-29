"""U0 end-to-end regression: the keyboard-only audit flow against a real server
on the WORKING DATABASE lane, where Review Decisions live (U0 exit criterion).

Self-contained and offline: seeds one Economy's Corpus from the committed
fixture Documents, runs it on the fake Engine (no key, no network, no model
call), spawns `regcompass serve` on that database, drives the pre-built UI with
Playwright (chromium), and tears down. Dev machine only (needs playwright + its
browser); never on the judge path.

This lane is the one that takes decisions: A / R / F write to the `reviews`
table, keyed by (Run, Mapping), and the flow below proves the decision and its
note survive a page reload. The frozen bundle lane (`serve --bundle`) is a
reading room and answers 409 to a decision, so it is not driven here.

It drives the COMMITTED interface bundle, so run `npm run build` in `ui/` first
when the React sources have moved ahead of `src/regcompass/ui_dist/`: the note
field and the export preview below arrived with the review lane, and a stale
bundle has neither. The same decision lane is covered without a browser, through the
endpoints the interface calls, by tests/test_reviews.py.

Run: uv run python scripts/u0_e2e.py
"""

from __future__ import annotations

import io
import json
import re
import subprocess
import sys
import tempfile
import time
import urllib.request
from pathlib import Path

from openpyxl import load_workbook
from playwright.sync_api import Page, expect, sync_playwright

from regcompass.workbook import SHEET, WORKBOOK_COLUMNS

ROOT = Path(__file__).resolve().parents[1]
PORT = 8770
BASE = f"http://127.0.0.1:{PORT}"
NOTE = "checked against the printed act"

# The two Evidence lists, by the ids their rows carry.
DOC_ROWS = '[data-testid^="doc-row-"]'
QUEUE_ROWS = '[data-testid^="queue-row-"]'

# A Document added with no address, so the flow can record one on it.
NO_URL_PDF = (
    ROOT / "tests/fixtures/sample_legislation/born_digital"
    / "PERSONAL DATA PROTECTION ACT 2010.pdf"
)
NO_URL_TITLE = "Personal Data Protection Act 2010"
NO_URL_ADDRESS = (
    "https://lom.agc.gov.my/ilims/upload/portal/akta/outputaktap/Act%20709%20ori.pdf"
)


def seed_run(db: Path, data: Path) -> str:
    """One fake-Engine Run over the committed SG fixture Corpus, through the
    same ingest path Discovery uses. Returns the Run Record id.

    The AU Corpus is seeded too but NOT run. The flow below starts that Run
    through the interface's own endpoint, which is the only way to fill the
    server's progress-line buffer and to make a second Economy the newest
    completed Run, and both of those are things the screens must handle.
    """
    sys.path.insert(0, str(ROOT / "tests"))
    from corpus_fixtures import seed_corpus  # dev-only test data

    from regcompass.engines import resolve_engine
    from regcompass.pipeline import run_economy
    from regcompass.storage import Storage

    from regcompass.corpus import add_document

    storage = Storage(db)
    storage.apply_schema()
    seed_corpus(storage, data, "SG")
    seed_corpus(storage, data, "AU")
    # One Document uploaded the way a reviewer does against the clock: no
    # Source URL. The flow below records its address through the interface.
    add_document(
        storage, data, "MY", NO_URL_PDF.read_bytes(), source_url=None,
        language="English", filename_hint=NO_URL_PDF.name,
        title=NO_URL_TITLE,
    )
    report = run_economy(
        storage, "SG", pillars=(7,), engine=resolve_engine("fake"), data_dir=data
    )
    storage.close()
    return report.run_id


def wait_ready(timeout_s: float = 60.0) -> None:
    deadline = time.time() + timeout_s
    while time.time() < deadline:
        try:
            with urllib.request.urlopen(f"{BASE}/api/documents", timeout=2):
                return
        except OSError:
            time.sleep(0.5)
    raise RuntimeError("server never became ready")


def run_economy_through_the_server(economy: str, timeout_s: float = 300.0) -> None:
    """Start a fake-Engine Run the way the Run panel starts one, and wait for
    it. The interface offers only the two declared Engines, so the fake one is
    reached here through the endpoint underneath the button; everything the
    server does afterwards (the line buffer, the Run Record) is identical."""
    body = json.dumps(
        {"economy": economy, "pillars": [7], "engine": "fake", "indicators": None}
    ).encode()
    request = urllib.request.Request(
        f"{BASE}/api/run", data=body, headers={"Content-Type": "application/json"}
    )
    with urllib.request.urlopen(request, timeout=30) as response:
        assert response.status == 200, response.status
    deadline = time.time() + timeout_s
    while time.time() < deadline:
        with urllib.request.urlopen(f"{BASE}/api/status", timeout=5) as response:
            state = json.load(response)
        if state["status"] == "done":
            return
        if state["status"] == "error":
            raise RuntimeError(f"the {economy} Run failed: {state.get('error')}")
        time.sleep(0.5)
    raise RuntimeError(f"the {economy} Run never finished")


def decide(page: Page, act) -> dict:
    """One Review Decision, taken and confirmed by the server.

    A decision is asynchronous: the interface posts it, and only when the
    response lands does it write the new status and advance to the next
    unreviewed record. Waiting for that response here is what keeps the next
    keypress from racing the advance (press A then K too quickly and K moves
    first, then the advance overrules it). Returns the stored Review Decision,
    so the caller can assert what the DATABASE holds and not merely what the
    screen says.
    """
    with page.expect_response(
        lambda r: r.request.method == "POST" and r.url.endswith("/api/reviews")
    ) as got:
        act()
    response = got.value
    assert response.status == 200, f"review POST {response.status}: {response.text()}"
    return response.json()


def choose_economy(page, code: str) -> None:
    """Pick an Economy on the Run panel: its card, opening "More Economies"
    first when the card is folded away there."""
    card = page.locator(f"label.pick-card:has(input[name=economy][value={code}])")
    if card.count() == 0:
        page.click("button.more-btn:has-text('More Economies')")
    card.first.click()


def run_flow(shots: Path, run_id: str) -> None:
    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page(viewport={"width": 1440, "height": 900})
        # Every request the page makes, so an idle screen can be checked for
        # the render loop below.
        requests: list[str] = []
        page.on("request", lambda r: requests.append(r.url))
        page.goto(BASE)

        # ------------------------------------------------------------------
        # The Run panel, before anything else: what a judge sees on opening.
        # ------------------------------------------------------------------

        # An idle screen is idle. A callback passed as a fresh arrow on every
        # render, read as an effect dependency, fetches on every render and
        # sets state, which renders again: a slow request loop that runs for
        # the life of the page. Measured at about one a second before the fix.
        page.wait_for_timeout(3000)
        corpus_calls = [u for u in requests if "/api/corpus" in u]
        assert len(corpus_calls) <= 3, (
            f"the Run panel re-fetched the Corpus {len(corpus_calls)} times while"
            " idle: an effect is looping"
        )

        # The Engine control opens on the REGISTRY's default, not on whichever
        # Engine carries a key. One click on Run must not spend on Engine A.
        default_engine = page.request.get(f"{BASE}/api/status").json()["default_engine"]
        assert default_engine, "the status endpoint names no default Engine"
        expect(page.locator("input[name=engine]:checked")).to_have_value(
            default_engine, timeout=20000
        )
        expect(page.get_by_test_id("engine-price")).to_contain_text(
            "USD per million tokens"
        )
        expect(page.get_by_test_id("engine-price")).to_contain_text("registry default")

        # An Economy with nothing in its Corpus says so, and names the three
        # ways out, rather than showing one button that refuses when pressed.
        choose_economy(page, "ID")
        empty = page.get_by_test_id("corpus-empty-state")
        expect(empty).to_be_visible(timeout=20000)
        expect(empty).to_contain_text("No Documents yet for Indonesia")
        expect(empty).to_contain_text("Discover")
        expect(empty).to_contain_text("Add document")
        expect(empty).to_contain_text("regcompass seed --economy ID")
        expect(page.get_by_role("button", name="Discover")).to_be_visible()
        page.screenshot(path=str(shots / "u0_empty_corpus.png"))

        # A seeded Economy says nothing of the kind.
        choose_economy(page, "SG")
        expect(empty).to_have_count(0, timeout=20000)

        # The Corpus of the chosen Economy is listed where a reviewer adds to
        # it, so an upload is visible without waiting for a Run.
        page.click("button.link-btn:has-text('Add document')")
        expect(page.get_by_test_id("corpus-table")).to_be_visible(timeout=20000)
        expect(page.locator("table.corpus tbody tr")).to_have_count(1)
        expect(page.locator("table.corpus tbody tr")).to_contain_text("Telecommunications")

        # ------------------------------------------------------------------
        # A Document added with no address can be given one HERE, on the row
        # that shows the problem. Re-adding it by URL would fetch the file
        # again and leave a second Document, so this is the only way back.
        # ------------------------------------------------------------------
        choose_economy(page, "MY")
        row = page.locator("table.corpus tbody tr", has_text=NO_URL_TITLE)
        expect(row).to_be_visible(timeout=20000)
        expect(row).to_contain_text("no Source URL")
        no_url_doc = page.get_by_test_id(re.compile(r"^set-source-url-")).first
        document_id = (no_url_doc.get_attribute("data-testid") or "").removeprefix(
            "set-source-url-"
        )
        assert document_id, "no Document is flagged as having no Source URL"
        no_url_doc.click()

        field = page.get_by_test_id(f"source-url-input-{document_id}")
        expect(field).to_be_visible()
        field.fill(NO_URL_ADDRESS)
        page.get_by_test_id(f"save-source-url-{document_id}").click()

        # The row refreshes and stops complaining, without a page reload.
        expect(row).not_to_contain_text("no Source URL", timeout=20000)
        expect(page.get_by_test_id(f"set-source-url-{document_id}")).to_have_count(0)
        page.screenshot(path=str(shots / "u0_source_url_recorded.png"))

        # And the server really holds it, not just the screen.
        listed = page.request.get(f"{BASE}/api/corpus?economy=MY").json()
        recorded = [d for d in listed["documents"] if d["document_id"] == document_id]
        assert recorded and recorded[0]["source_url"] == NO_URL_ADDRESS, recorded

        choose_economy(page, "SG")

        # Nothing is selectable on this screen, so no select-and-open keys.
        expect(page.get_by_test_id("select-hint")).to_have_count(0)

        # the review flow lives on Evidence, where rows ARE selectable
        page.click("button.nav-btn:has-text('Evidence')")
        expect(page.get_by_test_id("select-hint")).to_be_visible()

        # The screen says WHICH Run it is showing, on a plain page load where
        # nobody named one: the Run id, its Economy, its Pillar, its Engine.
        context = page.get_by_test_id("context")
        expect(context).to_contain_text(run_id)
        expect(context).to_contain_text("SG")
        expect(context).to_contain_text("Pillar 7")
        expect(page.locator(DOC_ROWS)).to_have_count(1)
        page.screenshot(path=str(shots / "u0_list.png"))

        # keyboard only: Enter opens the focused document
        page.keyboard.press("Enter")
        expect(page.get_by_test_id("mapping-position")).to_contain_text("Mapping 1 of")

        # PDF renders with the quote highlight overlaid
        page.wait_for_selector('[data-testid="pdf-page"] canvas', timeout=20000)
        page.wait_for_selector('[data-testid="quote-highlight"]', timeout=20000)
        assert page.get_by_test_id("quote-highlight").count() > 0, "no highlight rects"

        # the row's way out to the source: a real anchor, new tab, and (this
        # fixture Document carries a Portal address) the official URL at the
        # cited page rather than the copy on disk
        link = page.locator('[data-testid="record-pane"] .source-link')
        expect(link).to_be_visible()
        expect(link).to_have_attribute("target", "_blank")
        expect(link).to_have_attribute("rel", "noopener noreferrer")
        href = link.get_attribute("href") or ""
        assert href.startswith("https://"), f"source link is not official: {href!r}"
        assert "#page=" in href, f"source link carries no page fragment: {href!r}"

        # The Confidence composite is READABLE, not hidden in a hover tooltip:
        # the dots for the glance, the number for the reader who has to check
        # it against the Evidence Export.
        value = page.get_by_test_id("confidence-value").first
        expect(value).to_be_visible()
        assert re.fullmatch(r"[01]\.\d\d", value.inner_text().strip()), (
            f"the Confidence must read as a number: {value.inner_text()!r}"
        )

        page.wait_for_timeout(600)
        page.screenshot(path=str(shots / "u0_audit_highlight.png"))

        # accept via keyboard; the flow advances to the next unreviewed record
        review = decide(page, lambda: page.keyboard.press("a"))
        assert review["review_status"] == "accepted", review
        expect(page.get_by_test_id("mapping-position")).to_contain_text("Mapping 2 of")

        # k goes back; record 1 shows its accepted state
        page.keyboard.press("k")
        expect(page.get_by_test_id("mapping-position")).to_contain_text("Mapping 1 of")
        expect(page.get_by_test_id("record-decision")).to_contain_text("Accepted")

        # a note, then flag by mouse: the letter keys type inside the field.
        # Mouse and keyboard take the same path, so this advances as well: the
        # stored decision below is the proof, and k returns to read the badge.
        page.fill("#review-note", NOTE)
        review = decide(page, lambda: page.get_by_test_id("decide-flag").click())
        assert review["review_status"] == "flagged", review
        assert review["comment"] == NOTE, review
        expect(page.get_by_test_id("mapping-position")).to_contain_text("Mapping 2 of")
        page.keyboard.press("k")
        expect(page.get_by_test_id("mapping-position")).to_contain_text("Mapping 1 of")
        expect(page.get_by_test_id("record-decision")).to_contain_text("Flagged")

        # the decision and its note survive a reload: they are in the database
        page.reload()
        page.click("button.nav-btn:has-text('Evidence')")
        # Wait for the list before the keypress: Enter opens the FOCUSED row,
        # and the row is only focusable once the documents have arrived.
        expect(page.locator(DOC_ROWS)).to_have_count(1)
        page.keyboard.press("Enter")
        expect(page.get_by_test_id("mapping-position")).to_contain_text("Mapping 1 of")
        expect(page.get_by_test_id("record-decision")).to_contain_text("Flagged")
        expect(page.locator("#review-note")).to_have_value(NOTE)

        # accept it again: the second decision replaces the first
        review = decide(page, lambda: page.keyboard.press("a"))
        assert review["review_status"] == "accepted", review
        expect(page.get_by_test_id("mapping-position")).to_contain_text("Mapping 2 of")
        page.keyboard.press("k")
        expect(page.get_by_test_id("mapping-position")).to_contain_text("Mapping 1 of")
        expect(page.get_by_test_id("record-decision")).to_contain_text("Accepted")
        page.screenshot(path=str(shots / "u0_record_accepted.png"))

        # Escape returns to the list; the export preview counts the decision
        page.keyboard.press("Escape")
        expect(page.get_by_test_id("export-preview")).to_contain_text("accepted, will export")
        expect(page.get_by_test_id("export-accepted-count")).to_have_text("1")

        # export from the UI; the bottom bar reports the gated result
        page.get_by_role("button", name="Export workbook").click()
        expect(page.get_by_test_id("export-note")).to_contain_text("exported", timeout=60000)
        note = page.get_by_test_id("export-note").inner_text()
        print("export note:", note)
        # Plain words: how much shipped, then what was held back.
        assert re.search(r"\b\d+ rows exported, 1 accepted\b", note), note
        assert "unreviewed" in note, note
        # The reviewer is told the export is ready, not where it sits on a
        # filesystem they cannot reach.
        assert "Ready to download" in note, note
        assert "/" not in note.split("exported", 1)[1], (
            f"the footer must not show a server path: {note!r}"
        )

        # ... and the files are LINKS. The workbook is the primary control.
        workbook = page.get_by_test_id("download-workbook")
        expect(workbook).to_be_visible()
        expect(workbook).to_have_text("Download workbook")
        href = workbook.get_attribute("href") or ""
        assert "/api/outputs/download?name=submission.xlsx" in href, href
        others = page.locator('[data-testid="export-downloads"] a.download-link')
        assert others.count() >= 1, "the companion files must be offered too"

        # Follow the link the interface drew: the bytes it serves are the
        # organizer's own sheet, with their header row.
        got = page.request.get(BASE + href)
        assert got.status == 200, got.status
        sheet = load_workbook(io.BytesIO(got.body()))[SHEET]
        header = tuple(
            sheet.cell(row=4, column=c).value for c in range(1, len(WORKBOOK_COLUMNS) + 1)
        )
        assert header == WORKBOOK_COLUMNS, header
        page.screenshot(path=str(shots / "u0_export_downloads.png"))

        # ------------------------------------------------------------------
        # The Review queue: the Run's records in the order a person should
        # read them, lowest Confidence first, whichever Document they sit in.
        # One record carries a decision at this point (accepted above), which
        # is what the Unreviewed only filter is measured against.
        # ------------------------------------------------------------------
        page.get_by_test_id("view-queue").click()
        counts = page.get_by_test_id("queue-counts")
        expect(counts).to_be_visible(timeout=20000)
        line = counts.inner_text().strip()
        assert re.fullmatch(
            r"\d+ of \d+ below [01]\.\d\d Confidence, \d+ not reviewed yet\.", line
        ), line

        rows = page.locator(QUEUE_ROWS)
        expect(rows.first).to_be_visible(timeout=20000)
        listed = rows.count()
        texts = [
            t.strip()
            for t in page.locator('[data-testid="queue-confidence"]').all_inner_texts()
        ]
        assert len(texts) == listed, (len(texts), listed)
        scores = [float(t) for t in texts if re.fullmatch(r"[01]\.\d\d", t)]
        assert len(scores) == listed > 1, (scores, listed)
        assert scores == sorted(scores), f"the queue is not lowest first: {scores}"
        assert scores[0] == min(scores)
        page.screenshot(path=str(shots / "u0_review_queue.png"))

        # The filter hides the rows that carry a decision: exactly the one
        # accepted above disappears, and what is left is all undecided.
        page.get_by_test_id("queue-unreviewed-only").check()
        expect(rows).to_have_count(listed - 1, timeout=20000)
        states = [
            s.strip()
            for s in page.locator(f'{QUEUE_ROWS} [data-testid="queue-decision"]').all_inner_texts()
        ]
        assert len(states) == listed - 1, (len(states), listed)
        assert set(states) == {"Not reviewed"}, states
        # The count line still describes the whole Run, not the filtered view:
        # it says how much work there is, not how much is on screen.
        filtered_line = counts.inner_text().strip()
        assert filtered_line.split(" below ")[0] == line.split(" below ")[0], (
            filtered_line, line,
        )
        assert filtered_line.endswith(f"{listed - 1} not reviewed yet."), filtered_line

        page.get_by_test_id("queue-unreviewed-only").uncheck()
        expect(rows).to_have_count(listed, timeout=20000)

        # Opening a queue row lands in the same audit view, and the keys step
        # in QUEUE order: the Confidence on screen follows the queue's own
        # sequence, which is not the order the Documents table reads in.
        page.keyboard.press("Enter")
        nav = page.get_by_test_id("mapping-position")
        expect(nav).to_contain_text(f"Mapping 1 of {listed}")
        expect(page.locator("button.back-btn")).to_contain_text("Review queue")
        value = page.get_by_test_id("confidence-value").first
        expect(value).to_have_text(f"{scores[0]:.2f}")
        page.wait_for_selector('[data-testid="pdf-page"] canvas', timeout=20000)

        page.keyboard.press("j")
        expect(nav).to_contain_text(f"Mapping 2 of {listed}")
        expect(value).to_have_text(f"{scores[1]:.2f}")

        # A decision taken here advances along the queue as well.
        review = decide(page, lambda: page.keyboard.press("r"))
        assert review["review_status"] == "rejected", review
        landed = int(re.search(r"Mapping (\d+) of", nav.inner_text()).group(1))
        assert landed > 2, nav.inner_text()
        expect(value).to_have_text(f"{scores[landed - 1]:.2f}")

        # Escape goes back to the queue, and the queue has read the decision.
        page.keyboard.press("Escape")
        expect(counts).to_be_visible(timeout=20000)
        decided_row = page.get_by_test_id(f"queue-row-{review['mapping_id']}")
        expect(decided_row).to_contain_text("Rejected")
        page.get_by_test_id("view-documents").click()
        expect(page.locator(DOC_ROWS)).to_have_count(1)

        # ------------------------------------------------------------------
        # A SECOND Run, started through the endpoint the Run panel posts to.
        # It makes Australia the newest completed Run and fills the server's
        # progress-line buffer, which is what the next three checks read.
        # ------------------------------------------------------------------
        run_economy_through_the_server("AU")
        page.reload()

        # The Run panel shows the last Run again, finished, and its raw log
        # is kept behind the "Show raw log" toggle.
        expect(page.get_by_test_id("run-view")).to_contain_text("Finished", timeout=20000)
        page.get_by_role("button", name="Show raw log").click()
        log = page.get_by_test_id("run-log")
        expect(log).to_contain_text("M1 extract |", timeout=20000)
        # A stage with nothing to do says so, under its own plain name, so the
        # numbers never appear to skip.
        expect(log).to_contain_text("M2 ocr |")
        expect(log).to_contain_text("skipped")
        page.screenshot(path=str(shots / "u0_run_log_kept.png"))

        # Clicking deeper into the Run: a Document, then one of its Mappings
        # on its page. Inside a Run the Mapping is shown to read: a stray A, R
        # or F (typed while reading the Piece) writes no Review Decision, and
        # reviewing is one link away, in Evidence.
        log.get_by_role("button", name="Close").click()
        page.locator(".rv-doc-open").first.click()
        expect(page.get_by_test_id("run-drill")).to_have_attribute("data-layer", "document")
        page.locator(".rvd-map").first.click()
        expect(page.get_by_test_id("run-drill")).to_have_attribute("data-layer", "mapping")
        expect(page.get_by_test_id("record-pane")).to_be_visible(timeout=20000)
        expect(page.get_by_test_id("record-read-only")).to_be_visible()
        expect(page.get_by_test_id("open-in-evidence")).to_be_visible()
        expect(page.get_by_test_id("decide-accept")).to_have_count(0)
        def reviews_now():
            with urllib.request.urlopen(f"{BASE}/api/reviews", timeout=5) as response:
                return json.load(response)

        before = reviews_now()
        writes: list[str] = []
        page.on(
            "request",
            lambda r: writes.append(r.url) if r.method != "GET" else None,
        )
        page.locator(".rvd-piece-text").first.click()
        for key in ("a", "r", "f", "j", "k"):
            page.keyboard.press(key)
        page.locator("body").press("a")
        page.wait_for_timeout(800)
        assert writes == [], f"a key press in the drill-down wrote: {writes}"
        assert reviews_now() == before, "a key press in the drill-down changed a Review Decision"
        page.screenshot(path=str(shots / "u0_drill_mapping.png"))
        page.keyboard.press("Escape")
        expect(page.get_by_test_id("run-drill")).to_have_count(0)
        expect(page.get_by_test_id("run-view")).to_contain_text("Finished")
        print("drill-down: key presses wrote nothing")

        # The Runs table either fits the window or says that it scrolls.
        page.click("button.nav-btn:has-text('Run history')")
        expect(page.locator("table.runs tr.row")).to_have_count(2)
        expect(page.get_by_test_id("select-hint")).to_be_visible()
        overflow = page.evaluate(
            "() => { const e = document.querySelector('.table-scroll');"
            " return e.scrollWidth - e.clientWidth; }"
        )
        if overflow > 1:
            expect(page.get_by_test_id("runs-scroll-cue")).to_be_visible()
        print(f"runs table overflow at 1440px: {overflow}px")

        # How long the Run took and at what concurrency: both on the row, both
        # read off the Run Record rather than guessed.
        # lower-cased: the stylesheet renders these headers in capitals
        headers = [h.strip().lower() for h in page.locator("table.runs thead th").all_inner_texts()]
        assert "took" in headers, headers
        assert "at once" in headers, headers
        assert "ended" not in headers, "Took replaces Ended, it does not join it"
        run_row = page.locator("table.runs tr.row", has_text="Australia").first
        cells = run_row.locator("td").all_inner_texts()
        took = cells[headers.index("took")].strip()
        conc = cells[headers.index("at once")].strip()
        # a non-breaking space keeps each number with its unit
        assert re.fullmatch(r"\d+(\.\d+)?\u00a0s|\d+\u00a0min \d+\u00a0s", took), took
        assert conc.isdigit(), f"concurrency cell: {conc!r}"
        page.screenshot(path=str(shots / "u0_runs_table.png"))

        # Comparison opens on the Run just made, not on the first Economy in
        # the registry, and its empty state says what to do next.
        page.click("button.nav-btn:has-text('Comparison')")
        expect(page.locator("#cmp-economy")).to_have_value("AU")
        expect(page.locator("#cmp-pillar")).to_have_value("7")
        expect(page.get_by_test_id("compare-empty")).to_contain_text("Compare")
        # Nothing selectable here either.
        expect(page.get_by_test_id("select-hint")).to_have_count(0)
        page.get_by_role("button", name="Compare", exact=True).click()
        expect(page.locator(".error")).to_contain_text("Run Australia Pillar 7")
        expect(page.locator(".error")).to_contain_text("other Engine")

        # Evidence names the Run it fell back to, which is now the AU one.
        page.click("button.nav-btn:has-text('Evidence')")
        expect(page.get_by_test_id("context")).to_contain_text("AU")
        expect(page.get_by_test_id("context")).not_to_contain_text(run_id)

        browser.close()


def main() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)
        db = tmp_path / "regcompass.db"
        data = tmp_path / "data"
        started = time.monotonic()
        run_id = seed_run(db, data)
        print(f"seeded Run {run_id} in {time.monotonic() - started:.0f}s")
        server = subprocess.Popen(
            [
                sys.executable,
                "-m",
                "regcompass.cli",
                "serve",
                "--db",
                str(db),
                "--data-dir",
                str(data),
                "--out",
                str(tmp_path / "out"),
                "--port",
                str(PORT),
            ],
            cwd=ROOT,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.STDOUT,
        )
        try:
            wait_ready()
            shots = tmp_path / "shots"
            shots.mkdir()
            run_flow(shots, run_id)
        finally:
            server.terminate()
            server.wait(timeout=10)
    print("U0 E2E OK")


if __name__ == "__main__":
    main()
