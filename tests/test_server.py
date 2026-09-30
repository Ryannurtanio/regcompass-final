"""Tests for the one RegCompass server (regcompass.server).

No network, no paid model: every Run goes through the fake Engine over a
Corpus seeded from the committed fixture Documents, and the database endpoints
run against a scratch SQLite DB built with the real schema. The FastAPI
TestClient talks to the app in-process; the interface itself is covered by its
API contract plus one served-page check, because the committed bundle is what
a judge runs and no JavaScript test runner is on that path.

The scripted narration lane the retired one-page demo had is gone: a server
test that proved nothing but a hard-coded string list proved nothing about the
pipeline.
"""

from __future__ import annotations

import json
import re
import time
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from regcompass.export import REVIEW_CONFIDENCE_THRESHOLD
from regcompass.server import MAX_LIMIT, create_app
from regcompass.storage import Storage

from forbidden_portal import FORBIDDEN, forbidden_config, point_loaders_at

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture()
def scratch_db(tmp_path):
    """A real, empty schema DB in a scratch dir (never data/)."""
    db = tmp_path / "web.db"
    storage = Storage(db)
    storage.apply_schema()
    storage.conn.close()
    return db


@pytest.fixture()
def client(scratch_db, tmp_path):
    out = tmp_path / "out"
    out.mkdir()
    app = create_app(
        db_path=scratch_db, out_dir=out, data_dir=tmp_path / "data",
        ui_dir=None,
    )
    return TestClient(app)


def _seeded_app(tmp_path, economy="SG", **kwargs):
    """An app whose db already holds a Corpus, built through the ingest path."""
    import sys

    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from corpus_fixtures import seed_corpus

    db = tmp_path / "seeded.db"
    data = tmp_path / "data"
    out = tmp_path / "out"
    out.mkdir(exist_ok=True)
    storage = Storage(db)
    storage.apply_schema()
    seed_corpus(storage, data, economy)
    storage.conn.close()
    app = create_app(
        db_path=db, out_dir=out, data_dir=data,
        ui_dir=None, **kwargs,
    )
    return app, db, data


def _wait_done(client, timeout=5.0):
    deadline = time.time() + timeout
    while time.time() < deadline:
        st = client.get("/api/status").json()
        if not st["active"] and st["status"] in ("done", "error"):
            return st
        time.sleep(0.02)
    raise AssertionError("the job did not finish in time")


def _replayed_lines(client) -> list[str]:
    """Every narration line of the finished job, off the SSE stream."""
    msgs: list[str] = []
    with client.stream("GET", "/api/events") as resp:
        for line in resp.iter_lines():
            if not line.startswith("data: "):
                continue
            try:
                obj = json.loads(line[len("data: ") :])
            except json.JSONDecodeError:
                continue
            if "msg" in obj:
                msgs.append(obj["msg"])
    return msgs


# ---------------------------------------------------------------------------
# Box 1: one app, one port; the one-page demo and its route are gone
# ---------------------------------------------------------------------------


class TestOneApp:
    def _with_ui(self, tmp_path, scratch_db):
        ui = tmp_path / "dist"
        (ui / "assets").mkdir(parents=True)
        (ui / "index.html").write_text("<!doctype html><title>RegCompass</title>")
        (ui / "assets" / "app.js").write_text("// bundle")
        out = tmp_path / "out"
        out.mkdir(exist_ok=True)
        return TestClient(
            create_app(
                db_path=scratch_db, out_dir=out, data_dir=tmp_path / "data",
                ui_dir=ui,
            )
        )

    def test_the_interface_is_served_at_the_root(self, tmp_path, scratch_db):
        c = self._with_ui(tmp_path, scratch_db)
        r = c.get("/")
        assert r.status_code == 200
        assert "RegCompass" in r.text
        # a client-side route falls back to index.html, not a 404
        assert "RegCompass" in c.get("/runs").text

    def test_the_api_is_not_shadowed_by_the_interface(self, tmp_path, scratch_db):
        c = self._with_ui(tmp_path, scratch_db)
        assert c.get("/api/status").status_code == 200
        assert c.get("/api/documents").status_code == 200
        assert c.get("/assets/app.js").status_code == 200
        assert c.get("/assets/typo.js").status_code == 404

    def test_the_one_page_demo_and_its_module_are_retired(self, client):
        assert client.get("/demo").status_code == 404
        assert client.get("/api/db/mappings").status_code == 200, "the API survived"
        with pytest.raises(ImportError):
            import regcompass.webdemo  # noqa: F401
        with pytest.raises(ImportError):
            import regcompass.audit_api  # noqa: F401
        assert not (ROOT / "src" / "regcompass" / "webstatic").exists()

    def test_serve_is_the_one_entry_point(self):
        from typer.testing import CliRunner

        from regcompass.cli import app as cli_app

        r = CliRunner().invoke(cli_app, ["serve", "--help"])
        assert r.exit_code == 0, r.output
        # Without terminal styling: a CI runner colours every flag, and the
        # codes fall between the dashes and the name.
        out = re.sub(r"\x1b\[[0-9;]*m", "", r.output)
        for flag in ("--db", "--data-dir", "--port", "--bundle"):
            assert flag in out

    def test_demo_and_audit_are_deprecated_aliases_that_route_to_serve(
        self, tmp_path, monkeypatch
    ):
        import uvicorn
        from typer.testing import CliRunner

        import regcompass.server as server_mod
        from regcompass.cli import app as cli_app

        seen: dict = {}
        monkeypatch.setattr(uvicorn, "run", lambda app_, **kw: seen.update(served=True))
        real_create = server_mod.create_app

        def spy(**kwargs):
            seen["kwargs"] = kwargs
            return real_create(**{**kwargs, "ui_dir": None})

        monkeypatch.setattr(server_mod, "create_app", spy)

        r = CliRunner().invoke(cli_app, ["demo", "--db", str(tmp_path / "d.db")])
        assert r.exit_code == 0, r.output
        assert "serve" in r.output and "note:" in r.output
        assert seen["served"] is True
        assert seen["kwargs"]["bundle_manifest"] is None

        bundle = ROOT / "audit_bundle"
        seen.clear()
        r2 = CliRunner().invoke(cli_app, ["audit", "--bundle", str(bundle)])
        assert r2.exit_code == 0, r2.output
        assert "serve --bundle" in r2.output
        assert seen["kwargs"]["bundle_manifest"] == bundle / "manifest.json"


# ---------------------------------------------------------------------------
# Box 2: a Run through the API, on the fake Engine, with its Run Record
# ---------------------------------------------------------------------------


class TestARunThroughTheApi:
    def test_a_fake_engine_run_streams_progress_and_files_a_run_record(self, tmp_path):
        app, db, _ = _seeded_app(tmp_path)
        c = TestClient(app)

        started = c.post(
            "/api/run", json={"economy": "SG", "pillars": [7], "engine": "fake"}
        )
        assert started.status_code == 200, started.text
        assert started.json()["engine"] == "fake"
        assert started.json()["corpus_documents"] == 1

        st = _wait_done(c, timeout=180.0)
        assert st["status"] == "done", st

        lines = _replayed_lines(c)
        assert any(line.startswith("M0 start") for line in lines)
        assert any(line.startswith("M9 done") for line in lines)

        listed = c.get("/api/runs").json()["runs"]
        assert listed, "the Run Record survives the Run"
        record = listed[0]
        assert record["economy"] == "SG"
        assert record["engine"] == "fake"
        assert record["status"] == "completed"
        # A Run never fetches: Documents fetched is Discovery's number.
        assert record["documents_fetched"] == 0

        one = c.get(f"/api/runs/{record['run_id']}")
        assert one.status_code == 200
        assert one.json()["run_id"] == record["run_id"]
        assert c.get("/api/runs/nope").status_code == 404

        dl = c.get(f"/api/runs/{record['run_id']}/download")
        assert dl.status_code == 200
        assert json.loads(dl.text)["run_id"] == record["run_id"]

    def test_a_run_on_an_empty_corpus_is_refused_with_409_and_an_explanation(
        self, client
    ):
        r = client.post("/api/run", json={"economy": "SG", "pillars": [7], "engine": "fake"})
        assert r.status_code == 409
        body = r.json()["detail"]
        assert body["corpus_empty"] is True
        assert body["economy"] == "SG"
        assert "Discovery" in body["message"] and "Singapore" in body["message"]
        assert client.get("/api/status").json()["status"] == "idle", "nothing was started"

    def test_bad_requests_are_refused_before_anything_starts(self, client):
        assert client.post("/api/run", json={"economy": "XX", "engine": "fake"}).status_code == 400
        assert (
            client.post(
                "/api/run", json={"economy": "SG", "pillars": [99], "engine": "fake"}
            ).status_code
            == 400
        )
        assert client.post("/api/run", json={"economy": "SG", "engine": "wizard"}).status_code == 400
        assert (
            client.post(
                "/api/run", json={"economy": "SG", "engine": "fake", "mode": "wat"}
            ).status_code
            == 400
        )
        narrowed = client.post(
            "/api/run",
            json={"economy": "SG", "pillars": [7], "engine": "fake", "indicators": ["6.1"]},
        )
        assert narrowed.status_code == 400
        assert "6.1" in narrowed.text and "7.1" in narrowed.text


# ---------------------------------------------------------------------------
# Box 3: the interface's controls come from the registries, not from markup
# ---------------------------------------------------------------------------


class TestControlsComeFromTheRegistry:
    def test_status_lists_every_configured_economy_and_pillar(self, client, scratch_db):
        from regcompass.config import load_pillars, load_portals

        body = client.get("/api/status").json()
        portals = load_portals()
        assert body["economies"] == list(portals)
        assert body["economy_names"] == {c: p.official_name for c, p in portals.items()}
        assert body["pillars"] == sorted(load_pillars())
        assert body["default_pillars"] == [6, 7]
        assert str(scratch_db) == body["db"]

    def test_indicators_of_one_pillar(self, client):
        body = client.get("/api/indicators", params={"pillar": 7}).json()
        assert body["pillar"] == 7
        assert [i["id"] for i in body["indicators"]] == ["7.1", "7.2", "7.3", "7.4", "7.5"]
        assert all(i["name"] for i in body["indicators"])
        assert client.get("/api/indicators", params={"pillar": 99}).status_code == 400

    def test_the_engine_control_reads_the_registry_and_never_a_key(self, client):
        body = client.get("/api/engines").json()
        by_name = {e["name"]: e for e in body["engines"]}
        assert {"engine-a", "engine-b", "fake"} <= set(by_name)
        assert by_name["engine-b"]["open_weights"] is True
        assert by_name["engine-a"]["display_name"]
        assert by_name["engine-a"]["usd_per_million_input_tokens"] > 0
        assert by_name["fake"]["key_set"] is True, "an Engine needing no key can always run"
        assert set(by_name["engine-a"]) == {
            "name", "display_name", "open_weights",
            "usd_per_million_input_tokens", "usd_per_million_output_tokens",
            "api_key_env", "key_set", "default",
        }

    def test_a_new_economy_appears_and_runs_without_a_bundle_rebuild(
        self, tmp_path, monkeypatch
    ):
        """The interface reads its Economy list at request time, so adding one
        to the registry needs no code edit and no rebuilt bundle."""
        import shutil

        import yaml

        from regcompass.config import CONFIG_DIR
        import regcompass.config as config_mod

        config_dir = tmp_path / "config"
        shutil.copytree(CONFIG_DIR, config_dir)
        portals = yaml.safe_load((config_dir / "portals.yaml").read_text(encoding="utf-8"))
        portals["portals"]["ZZ"] = dict(
            portals["portals"]["SG"], official_name="Testland"
        )
        (config_dir / "portals.yaml").write_text(yaml.safe_dump(portals), encoding="utf-8")
        monkeypatch.setattr(config_mod, "CONFIG_DIR", config_dir)
        config_mod.load_portals.cache_clear() if hasattr(
            config_mod.load_portals, "cache_clear"
        ) else None

        db = tmp_path / "zz.db"
        storage = Storage(db)
        storage.apply_schema()
        # A Corpus of one Document, so the Economy is not merely listed: a Run
        # on it is accepted and starts.
        storage.upsert_document(
            "doc_zz_testland_act", "ZZ", "z" * 64,
            full_text="A provision of the Testland Act about spectrum licences.",
            title="Testland Act",
        )
        storage.conn.close()
        out = tmp_path / "out"
        out.mkdir()
        c = TestClient(
            create_app(
                db_path=db, out_dir=out, data_dir=tmp_path / "data",
                ui_dir=None,
            )
        )
        body = c.get("/api/status").json()
        assert "ZZ" in body["economies"]
        assert body["economy_names"]["ZZ"] == "Testland"
        started = c.post("/api/run", json={"economy": "ZZ", "pillars": [7], "engine": "fake"})
        assert started.status_code == 200, started.text
        assert started.json()["economy"] == "ZZ"
        _wait_done(c, timeout=180.0)
        bad = c.post("/api/run", json={"economy": "QQ", "engine": "fake"})
        assert bad.status_code == 400
        assert "ZZ" in bad.json()["detail"]


# ---------------------------------------------------------------------------
# Box 4: the Settings key lives in memory and nowhere else
# ---------------------------------------------------------------------------


MARKER = "sk-regcompass-test-marker-3f9c1d7b"


class TestSettingsKey:
    @pytest.fixture(autouse=True)
    def _forget(self):
        from regcompass import engines as engines_mod

        engines_mod._SESSION_KEYS.clear()
        yield
        engines_mod._SESSION_KEYS.clear()

    def test_a_posted_key_enables_the_engine_and_leaves_no_trace(
        self, tmp_path, monkeypatch, caplog
    ):
        import logging

        monkeypatch.delenv("OPENROUTER_API_KEY", raising=False)
        app, db, data = _seeded_app(tmp_path)
        c = TestClient(app)

        assert {e["name"]: e["key_set"] for e in c.get("/api/engines").json()["engines"]}[
            "engine-a"
        ] is False
        refused = c.post("/api/run", json={"economy": "SG", "pillars": [7], "engine": "engine-a"})
        assert refused.status_code == 400
        assert "OPENROUTER_API_KEY" in refused.json()["detail"]

        caplog.set_level(logging.DEBUG)
        posted = c.post("/api/settings/key", json={"engine": "engine-a", "key": MARKER})
        assert posted.status_code == 200, posted.text
        assert posted.json() == {
            "engine": "engine-a", "env": "OPENROUTER_API_KEY", "key_set": True
        }
        assert MARKER not in posted.text, "the key is never echoed back"

        engines = {e["name"]: e for e in c.get("/api/engines").json()["engines"]}
        assert engines["engine-a"]["key_set"] is True
        assert engines["engine-b"]["key_set"] is True, "the same variable feeds both"

        # The key is held here, not exported into the process environment.
        import os

        assert os.environ.get("OPENROUTER_API_KEY") is None

        # Run the fake Engine while the key is held, then sweep everything the
        # server said, logged or wrote.
        started = c.post("/api/run", json={"economy": "SG", "pillars": [7], "engine": "fake"})
        assert started.status_code == 200, started.text
        _wait_done(c, timeout=180.0)

        bodies = [
            c.get("/api/status").text,
            c.get("/api/stats").text,
            c.get("/api/engines").text,
            c.get("/api/runs").text,
            c.get("/api/documents").text,
            c.get("/api/db/mappings").text,
            c.get("/api/outputs").text,
            "\n".join(_replayed_lines(c)),
        ]
        for body in bodies:
            assert MARKER not in body
        assert MARKER not in caplog.text

        for root in (data, tmp_path / "out", Path(db).parent):
            for path in Path(root).rglob("*"):
                if not path.is_file():
                    continue
                if MARKER.encode() in path.read_bytes():
                    raise AssertionError(f"the key reached {path}")

        # Restart: a fresh app on the same database has forgotten it.
        from regcompass import engines as engines_mod

        engines_mod._SESSION_KEYS.clear()
        restarted = TestClient(
            create_app(
                db_path=db, out_dir=tmp_path / "out", data_dir=data,
                ui_dir=None,
            )
        )
        after = {e["name"]: e for e in restarted.get("/api/engines").json()["engines"]}
        assert after["engine-a"]["key_set"] is False

    def test_a_key_can_be_forgotten_on_request(self, client, monkeypatch):
        monkeypatch.delenv("OPENROUTER_API_KEY", raising=False)
        client.post("/api/settings/key", json={"engine": "engine-b", "key": MARKER})
        engines = {e["name"]: e for e in client.get("/api/engines").json()["engines"]}
        assert engines["engine-b"]["key_set"] is True
        dropped = client.delete("/api/settings/key/engine-b")
        assert dropped.status_code == 200
        assert dropped.json()["key_set"] is False
        engines = {e["name"]: e for e in client.get("/api/engines").json()["engines"]}
        assert engines["engine-b"]["key_set"] is False

    def test_a_key_for_no_engine_at_all_is_a_400(self, client):
        assert client.post("/api/settings/key", json={"key": MARKER}).status_code == 400
        assert (
            client.post(
                "/api/settings/key", json={"engine": "wizard", "key": MARKER}
            ).status_code
            == 400
        )
        assert (
            client.post(
                "/api/settings/key", json={"engine": "fake", "key": MARKER}
            ).status_code
            == 400
        ), "the fake Engine needs no key"
        assert (
            client.post(
                "/api/settings/key", json={"env": "OPENROUTER_API_KEY", "key": "  "}
            ).status_code
            == 400
        )


# ---------------------------------------------------------------------------
# Box 5: the audit view reads a Run out of the working database
# ---------------------------------------------------------------------------


@pytest.fixture(scope="module")
def ran(tmp_path_factory):
    """One fake-Engine Run over the seeded SG Corpus, reused by every test that
    reads it back: a real Run through the server is the slowest thing here."""
    tmp_path = tmp_path_factory.mktemp("audit_db")
    app, db, data = _seeded_app(tmp_path)
    c = TestClient(app)
    started = c.post("/api/run", json={"economy": "SG", "pillars": [7], "engine": "fake"})
    assert started.status_code == 200, started.text
    st = _wait_done(c, timeout=240.0)
    assert st["status"] == "done", st
    run_id = c.get("/api/runs").json()["runs"][0]["run_id"]
    return c, run_id, tmp_path


class TestTheEvidenceScreenCanNameItsRun:
    """The audit reads fall back to the newest completed Run when nobody names
    one, and that fallback used to be invisible: a plain page load showed rows
    belonging to a Run the screen never mentioned."""

    def test_the_resolved_run_comes_back_with_its_economy_pillar_and_engine(self, ran):
        c, run_id, _ = ran
        named = c.get("/api/audit/run", params={"run_id": run_id}).json()
        assert named["run_id"] == run_id
        record = named["record"]
        assert record["economy"] == "SG"
        assert record["pillars"] == [7]
        assert record["engine"] == "fake"
        assert named["bundle_mode"] is False

        # No Run named: the SAME Run the documents endpoint falls back to.
        fallback = c.get("/api/audit/run").json()
        assert fallback["run_id"] == run_id
        assert fallback["record"]["run_id"] == run_id

    def test_a_database_with_no_run_answers_honestly_rather_than_failing(self, client):
        body = client.get("/api/audit/run").json()
        assert body["run_id"] is None
        assert body["record"] is None


class TestTheStatusNamesTheDefaultEngine:
    """The Run panel opens on the Engine the registry declares, not on
    whichever Engine happens to carry a key. The two declared Engines are not
    interchangeable: one is open-weight and one is billed per call, so a Run
    panel that preselects the billed one turns a single click into spending
    nobody chose."""

    def test_the_status_carries_the_registrys_default_engine(self, client):
        from regcompass.config import load_models

        declared = load_models().default_engine
        body = client.get("/api/status").json()
        assert body["default_engine"] == declared

    def test_it_agrees_with_the_engine_registry_and_names_a_real_engine(
        self, client
    ):
        status = client.get("/api/status").json()
        registry = client.get("/api/engines").json()
        assert status["default_engine"] == registry["default_engine"]
        by_name = {e["name"]: e for e in registry["engines"]}
        assert status["default_engine"] in by_name
        assert by_name[status["default_engine"]]["default"] is True


class TestTheProgressLogOutlivesTheRun:
    def test_an_idle_server_holds_no_lines(self, client):
        body = client.get("/api/run/log").json()
        assert body["lines"] == []
        assert body["status"] == "idle"
        assert body["active"] is False

    def test_the_finished_runs_lines_are_still_there_to_read_back(self, ran):
        """On completion the app moves the reviewer to the evidence and the Run
        panel unmounts. Coming back to a placeholder reads as a Run that never
        happened, so the panel asks the server for the lines it still holds."""
        c, run_id, _ = ran
        body = c.get("/api/run/log").json()
        assert body["status"] == "done"
        assert body["active"] is False
        assert body["run_id"] == run_id
        assert body["lines"], "the finished Run's narration must still be readable"
        assert any(line.startswith("M1 extract |") for line in body["lines"])


class TestAuditViewOverTheWorkingDatabase:
    def test_the_run_s_documents_and_mappings_come_back(self, ran):
        c, run_id, _ = ran
        docs = c.get("/api/documents", params={"run_id": run_id}).json()
        assert docs, "the Run produced evidence"
        doc = docs[0]
        assert doc["document_id"].startswith("doc_sg_")
        assert doc["economy"] == "SG"
        assert doc["n_records"] > 0
        assert doc["n_pages"] > 0

        recs = c.get(
            f"/api/documents/{doc['document_id']}/records", params={"run_id": run_id}
        ).json()
        assert recs
        assert all(r["indicator_id"].startswith("7.") for r in recs)
        assert all(r["quote_preview"] for r in recs)
        assert any(r["confidence"] is not None for r in recs)

    def test_each_scored_record_carries_parts_that_add_up_to_its_confidence(self, ran):
        c, run_id, _ = ran
        docs = c.get("/api/documents", params={"run_id": run_id}).json()
        recs = c.get(
            f"/api/documents/{docs[0]['document_id']}/records", params={"run_id": run_id}
        ).json()
        scored = [r for r in recs if r["confidence"] is not None]
        assert scored
        for r in scored:
            parts = r["confidence_parts"]
            assert [p["signal"] for p in parts] == [
                "similarity", "quote_length", "specificity", "attempts",
            ]
            assert round(sum(p["contribution"] for p in parts), 2) == r["confidence"]
        # The Review queue is built from the same rows, so it carries them too.
        queue = c.get("/api/records", params={"run_id": run_id}).json()["records"]
        for r in queue:
            if r["confidence"] is not None:
                assert round(sum(p["contribution"] for p in r["confidence_parts"]), 2) == r["confidence"]

    def test_the_quote_is_highlighted_from_the_stored_word_boxes(self, ran):
        c, run_id, _ = ran
        docs = c.get("/api/documents", params={"run_id": run_id}).json()
        recs = c.get(
            f"/api/documents/{docs[0]['document_id']}/records", params={"run_id": run_id}
        ).json()
        detail = c.get(
            f"/api/records/{recs[0]['mapping_id']}", params={"run_id": run_id}
        ).json()
        assert detail["highlight_available"] is True
        assert detail["highlights"], "rectangles, not an empty list"
        first = detail["highlights"][0]
        assert first["page"] >= 1
        assert first["x1"] > first["x0"] and first["y1"] > first["y0"]
        assert detail["record"]["verbatim_quote"]
        assert detail["quote_char_start"] is not None

    def test_the_source_pdf_is_served_from_the_corpus(self, ran):
        c, run_id, _ = ran
        docs = c.get("/api/documents", params={"run_id": run_id}).json()
        r = c.get(f"/api/documents/{docs[0]['document_id']}/pdf", params={"run_id": run_id})
        assert r.status_code == 200
        assert r.headers["content-type"] == "application/pdf"
        assert r.content[:5] == b"%PDF-"

    def test_without_a_run_id_the_latest_completed_run_is_shown(self, ran):
        c, run_id, _ = ran
        assert c.get("/api/documents").json() == c.get(
            "/api/documents", params={"run_id": run_id}
        ).json()

    def test_a_review_round_trips_and_unknown_ids_404(self, ran):
        c, run_id, _ = ran
        docs = c.get("/api/documents", params={"run_id": run_id}).json()
        recs = c.get(
            f"/api/documents/{docs[0]['document_id']}/records", params={"run_id": run_id}
        ).json()
        mapping_id = recs[0]["mapping_id"]
        r = c.post(
            "/api/reviews",
            json={
                "run_id": run_id, "mapping_id": mapping_id,
                "review_status": "accepted", "reviewer": "ryan",
            },
        )
        assert r.status_code == 200, r.text
        assert r.json()["run_id"] == run_id
        again = c.get("/api/documents", params={"run_id": run_id}).json()
        assert again[0]["n_accepted"] == 1
        detail = c.get(f"/api/records/{mapping_id}", params={"run_id": run_id}).json()
        assert detail["review"]["review_status"] == "accepted"

        assert c.get("/api/documents/doc_nope/records").status_code == 404
        assert c.get("/api/records/nope").status_code == 404
        assert (
            c.post(
                "/api/reviews",
                json={"run_id": run_id, "mapping_id": "nope", "review_status": "accepted"},
            ).status_code
            == 404
        )

    def test_an_empty_database_lists_nothing_instead_of_failing(self, tmp_path):
        c = TestClient(
            create_app(
                db_path=tmp_path / "nothing.db", out_dir=tmp_path,
                data_dir=tmp_path / "data",
                ui_dir=None,
            )
        )
        assert c.get("/api/documents").json() == []
        assert c.get("/api/documents/doc_x/records").status_code == 404
        queue = c.get("/api/records").json()
        assert queue["records"] == []
        assert (queue["total"], queue["below_threshold"], queue["unreviewed"]) == (0, 0, 0)
        assert queue["threshold"] == REVIEW_CONFIDENCE_THRESHOLD


def _sort_key(row: dict) -> tuple[float, str]:
    """How the queue orders one row: confidence ascending, a row with no
    Confidence first because nothing is known about it, ties by Mapping id."""
    return (-1.0 if row["confidence"] is None else row["confidence"], row["mapping_id"])


class TestTheReviewQueue:
    """The Run's Mappings as one list, lowest Confidence first, so the rows the
    calibration rule says to hand-check are the rows a reviewer meets first.
    The Documents table answers "what did this Run find"; this answers "what do
    I have to look at"."""

    def test_the_queue_is_ordered_by_confidence_and_names_each_document(self, ran):
        c, run_id, _ = ran
        body = c.get("/api/records", params={"run_id": run_id}).json()
        rows = body["records"]
        assert rows, "the Run produced evidence"
        assert body["run_id"] == run_id
        assert body["total"] == len(rows)

        docs = c.get("/api/documents", params={"run_id": run_id}).json()
        assert body["total"] == sum(d["n_records"] for d in docs)
        titles = {d["document_id"]: d["title"] for d in docs}
        for row in rows:
            assert row["document_id"] in titles
            assert row["document_title"] == titles[row["document_id"]]
            # The queue row carries everything the Documents table's rows do,
            # so opening one lands in the same record pane.
            assert row["quote_preview"] and row["indicator_id"]

        assert [_sort_key(r) for r in rows] == sorted(_sort_key(r) for r in rows)

    def test_the_count_block_states_the_threshold_and_what_sits_below_it(self, ran):
        c, run_id, _ = ran
        body = c.get("/api/records", params={"run_id": run_id}).json()
        rows = body["records"]
        threshold = body["threshold"]
        assert threshold == REVIEW_CONFIDENCE_THRESHOLD
        below = [
            r for r in rows if r["confidence"] is None or r["confidence"] < threshold
        ]
        assert body["below_threshold"] == len(below)
        assert body["unreviewed"] == len([r for r in rows if r["review_status"] is None])
        assert body["below_threshold"] <= body["total"]

    def test_the_threshold_the_server_reports_is_the_number_the_readme_states(self):
        readme = (ROOT / "README.md").read_text(encoding="utf-8")
        assert f"below {REVIEW_CONFIDENCE_THRESHOLD:.2f} by hand" in readme

    def test_unreviewed_only_hides_the_rows_that_carry_a_decision(self, ran):
        c, run_id, _ = ran
        before = c.get("/api/records", params={"run_id": run_id}).json()
        # The highest-Confidence row, so the ordering the other tests read is
        # untouched by the decision this one takes.
        decided = next(
            r["mapping_id"]
            for r in reversed(before["records"])
            if r["review_status"] is None
        )
        assert c.post(
            "/api/reviews",
            json={
                "run_id": run_id, "mapping_id": decided,
                "review_status": "accepted",
            },
        ).status_code == 200

        filtered = c.get(
            "/api/records", params={"run_id": run_id, "unreviewed": True}
        ).json()
        ids = [r["mapping_id"] for r in filtered["records"]]
        assert decided not in ids
        assert all(r["review_status"] is None for r in filtered["records"])
        assert len(ids) == filtered["unreviewed"]
        # The count block counts the whole Run, not the filtered view: the line
        # on screen has to keep saying how many rows there are in all.
        assert filtered["total"] == before["total"]
        assert filtered["below_threshold"] == before["below_threshold"]
        assert filtered["unreviewed"] < before["unreviewed"]
        assert [_sort_key(r) for r in filtered["records"]] == sorted(
            _sort_key(r) for r in filtered["records"]
        )

    def test_a_row_with_no_confidence_sorts_first_and_counts_as_below(self):
        """Nothing is known about an unscored row, which is the strongest
        reason to put a person in front of it. Every Mapping of the fixture Run
        scores, so the rule is held here against a source built to carry one."""
        from regcompass.audit import DocumentSummary, RecordSummary, build_review_queue

        def record(mapping_id: str, confidence: float | None) -> RecordSummary:
            return RecordSummary(
                mapping_id=mapping_id, indicator_id="7.2",
                indicator_name="Lack of dedicated legal framework for cybersecurity",
                section="Part 2 s. 14", subsection=None, page_number=1,
                quote_preview="a verbatim quote", confidence=confidence,
                controlling_evidence=False, review_status=None,
            )

        class OneUnscoredRow:
            run_id = "run_x"

            def document_summaries(self, reviews):
                return [
                    DocumentSummary(
                        document_id="doc_a", title="An Act", economy="SG",
                        n_pages=1, ocr_applied=False, n_records=3,
                        n_accepted=0, n_rejected=0, n_flagged=0,
                    )
                ]

            def record_summaries(self, document_id, reviews):
                assert document_id == "doc_a"
                return [
                    record("m_high", 0.91),
                    record("m_unscored", None),
                    record("m_low", 0.41),
                ]

        queue = build_review_queue(OneUnscoredRow(), {})
        assert [r.mapping_id for r in queue.records] == ["m_unscored", "m_low", "m_high"]
        assert queue.records[0].confidence is None
        assert (queue.total, queue.unreviewed) == (3, 3)
        # The unscored row counts against the threshold with the low one.
        assert queue.below_threshold == 2

    def test_the_natural_order_is_still_reachable_and_a_bad_sort_is_refused(self, ran):
        c, run_id, _ = ran
        body = c.get("/api/records", params={"run_id": run_id, "sort": "record"}).json()
        rows = body["records"]
        keys = [(r["document_id"], r["indicator_id"], r["mapping_id"]) for r in rows]
        assert keys == sorted(keys)
        assert len(rows) == body["total"]
        assert c.get(
            "/api/records", params={"run_id": run_id, "sort": "wobble"}
        ).status_code == 400


class TestExportFromTheWorkingDatabase:
    """The Export button on the database lane. The review gate is the product
    rule (`only accepted Mappings enter the Evidence Export`), so it applies
    here exactly as it does on the frozen bundle: an unreviewed Run ships no
    substantive row, and the supplementary discloses why."""

    def test_the_review_gate_decides_what_ships(self, tmp_path):
        import csv as csv_mod

        from regcompass.export import ABSENCE_MARKER

        app, _, _ = _seeded_app(tmp_path)
        c = TestClient(app)
        assert c.post(
            "/api/run", json={"economy": "SG", "pillars": [7], "engine": "fake"}
        ).status_code == 200
        assert _wait_done(c, timeout=240.0)["status"] == "done"
        run_id = c.get("/api/runs").json()["runs"][0]["run_id"]

        docs = c.get("/api/documents", params={"run_id": run_id}).json()
        recs = c.get(
            f"/api/documents/{docs[0]['document_id']}/records", params={"run_id": run_id}
        ).json()
        assert len(recs) > 1, "the gate is only interesting with several Mappings"

        # Before any Review Decision: nothing a reviewer approved, so nothing
        # substantive ships, and every Mapping is disclosed as unreviewed.
        first = c.post("/api/export", params={"run_id": run_id})
        assert first.status_code == 200, first.text
        summary = first.json()
        assert summary["n_records_total"] == len(recs)
        assert summary["n_accepted"] == 0
        rows = list(csv_mod.DictReader(Path(summary["csv_path"]).open(encoding="utf-8-sig")))
        assert [r for r in rows if r["Article / Section"] != ABSENCE_MARKER] == []
        gate = json.loads(
            Path(summary["supplementary_path"]).read_text(encoding="utf-8")
        )["review_gate"]
        assert gate["n_verified"] == len(recs)
        assert gate["n_unreviewed"] == len(recs)
        assert gate["n_accepted"] == 0
        assert "accepted" in gate["rule"]

        # Accept exactly one: that Mapping, and only that one, ships. Not a 7.1
        # or 7.2 one: those ship once per Economy and only on a framework law,
        # and this fixture's Act is not one.
        accepted = next(r for r in recs if r["indicator_id"] not in ("7.1", "7.2"))
        other = next(r for r in recs if r["mapping_id"] != accepted["mapping_id"])
        assert c.post(
            "/api/reviews",
            json={
                "run_id": run_id, "mapping_id": accepted["mapping_id"],
                "review_status": "accepted",
            },
        ).status_code == 200
        c.post(
            "/api/reviews",
            json={
                "run_id": run_id, "mapping_id": other["mapping_id"],
                "review_status": "rejected",
            },
        )
        preview = c.get("/api/export/preview", params={"run_id": run_id}).json()
        assert preview["accepted_mapping_ids"] == [accepted["mapping_id"]]
        assert (preview["n_accepted"], preview["n_rejected"]) == (1, 1)
        second = c.post("/api/export", params={"run_id": run_id}).json()
        assert second["n_accepted"] == 1
        rows = list(csv_mod.DictReader(Path(second["csv_path"]).open(encoding="utf-8-sig")))
        substantive = [r for r in rows if r["Article / Section"] != ABSENCE_MARKER]
        assert len(substantive) == 1
        assert substantive[0]["Indicator ID"] == accepted["indicator_id"]
        # n_rows is a CSV-parsed count: verbatim quotes embed newlines.
        assert second["n_rows"] == len(rows)
        gate = json.loads(
            Path(second["supplementary_path"]).read_text(encoding="utf-8")
        )["review_gate"]
        assert (gate["n_accepted"], gate["n_rejected"]) == (1, 1)
        assert gate["n_unreviewed"] == len(recs) - 2

        # A correction is saved beside the accepted row, and the corrected
        # Mapping ships too, under the reviewer's Indicator.
        own = other["indicator_id"]
        detail = c.get(
            f"/api/records/{other['mapping_id']}", params={"run_id": run_id}
        ).json()
        to = next(ch["id"] for ch in detail["correction_choices"] if ch["id"] not in (own, "7.1", "7.2"))
        corrected = c.post(
            "/api/reviews",
            json={
                "run_id": run_id, "mapping_id": other["mapping_id"],
                "review_status": "corrected", "corrected_indicator_id": to,
                "comment": "Right provision, other Indicator.",
            },
        )
        assert corrected.status_code == 200, corrected.text
        third = c.post("/api/export", params={"run_id": run_id})
        assert third.status_code == 200, third.text
        assert third.json()["n_corrected"] == 1
        rows = list(csv_mod.DictReader(Path(third.json()["csv_path"]).open(encoding="utf-8-sig")))
        substantive = [r for r in rows if r["Article / Section"] != ABSENCE_MARKER]
        assert len(substantive) == 2
        assert to in {r["Indicator ID"] for r in substantive}
        gate = json.loads(
            Path(third.json()["supplementary_path"]).read_text(encoding="utf-8")
        )["review_gate"]
        assert [o["mapping_id"] for o in gate["overrides"]] == [other["mapping_id"]]


class TestBundleLaneStillWorks:
    """`serve --bundle audit_bundle` is the judge's keyless path: no key, no
    Run, no working database. It must keep serving the Round 1 bundle."""

    @pytest.fixture()
    def bundle_client(self, tmp_path):
        manifest = ROOT / "audit_bundle" / "manifest.json"
        if not manifest.is_file():  # pragma: no cover - the bundle is committed
            pytest.skip("no committed audit bundle")
        out = tmp_path / "out"
        out.mkdir()
        return TestClient(
            create_app(
                db_path=tmp_path / "unused.db", out_dir=out,
                data_dir=tmp_path / "data", bundle_manifest=manifest,
                ui_dir=None,
            )
        )

    def test_the_bundle_documents_and_highlights_are_served(self, bundle_client):
        docs = bundle_client.get("/api/documents").json()
        assert docs
        doc_id = docs[0]["document_id"]
        recs = bundle_client.get(f"/api/documents/{doc_id}/records").json()
        assert recs
        detail = bundle_client.get(f"/api/records/{recs[0]['mapping_id']}").json()
        assert detail["highlight_available"] is True
        assert detail["highlights"]

    def test_the_bundle_has_a_review_queue_too(self, bundle_client):
        """The keyless lane is the one a desk reviewer opens, so the queue has
        to be there as well. It takes no decisions (the bundle answers 409 to
        one), but it still orders the reading."""
        body = bundle_client.get("/api/records").json()
        rows = body["records"]
        assert rows
        assert body["total"] == len(rows)
        assert body["threshold"] == REVIEW_CONFIDENCE_THRESHOLD
        assert [_sort_key(r) for r in rows] == sorted(_sort_key(r) for r in rows)
        doc_ids = {d["document_id"] for d in bundle_client.get("/api/documents").json()}
        assert {r["document_id"] for r in rows} <= doc_ids
        assert body["unreviewed"] == body["total"], "a reading room holds no decisions"

    def test_the_bundle_pdf_is_served(self, bundle_client):
        docs = bundle_client.get("/api/documents").json()
        r = bundle_client.get(f"/api/documents/{docs[0]['document_id']}/pdf")
        assert r.status_code == 200
        assert r.content[:5] == b"%PDF-"

    def test_the_bundle_mode_says_so_on_status(self, bundle_client):
        assert bundle_client.get("/api/status").json()["bundle_mode"] is True

    def test_the_bundle_lane_takes_no_review_decision_and_exports_ungated(
        self, bundle_client, tmp_path
    ):
        """A Review Decision belongs to a Run in the working database; the
        frozen bundle has none, so it refuses one and its export ships what
        Round 1 shipped."""
        import csv as csv_mod

        docs = bundle_client.get("/api/documents").json()
        recs = bundle_client.get(f"/api/documents/{docs[0]['document_id']}/records").json()
        refused = bundle_client.post(
            "/api/reviews",
            json={"mapping_id": recs[0]["mapping_id"], "review_status": "accepted"},
        )
        assert refused.status_code == 409
        assert bundle_client.get("/api/export/preview").json()["gated"] is False
        summary = bundle_client.post("/api/export").json()
        rows = list(csv_mod.DictReader(Path(summary["csv_path"]).open(encoding="utf-8-sig")))
        assert summary["n_rows"] == len(rows)
        assert summary["n_records_total"] > 0

    def test_the_export_names_its_files_and_the_download_serves_each(
        self, bundle_client
    ):
        """What the reviewer takes away. The Export answers with the file NAMES
        the download endpoint takes, workbook first, so the interface can offer
        a link instead of a path nobody on the container path can reach."""
        import io

        from openpyxl import load_workbook

        from regcompass.workbook import SHEET, WORKBOOK_COLUMNS

        summary = bundle_client.post("/api/export").json()
        files = summary["files"]
        assert files, "the Export named no files"
        # Bare download names, never paths: the endpoint refuses anything else.
        for entry in files:
            assert "/" not in entry["name"], entry
            assert entry["label"].strip(), entry
        assert [f["name"] for f in files][:2] == ["submission.xlsx", "submission.csv"]
        assert [f["name"] for f in files if f["primary"]] == ["submission.xlsx"]
        assert {"submission.json", "supplementary.json"} <= {f["name"] for f in files}

        media = {
            "submission.xlsx": (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
            "submission.csv": "text/csv",
            "submission.json": "application/json",
            "supplementary.json": "application/json",
        }
        for entry in files:
            got = bundle_client.get(
                "/api/outputs/download", params={"name": entry["name"]}
            )
            assert got.status_code == 200, entry["name"]
            assert got.headers["content-type"].split(";")[0] == media[entry["name"]]
            assert got.content, entry["name"]

        # The workbook the link streams is the organizer's own sheet, not an
        # empty file with the right name.
        workbook = bundle_client.get(
            "/api/outputs/download", params={"name": "submission.xlsx"}
        )
        sheet = load_workbook(io.BytesIO(workbook.content))[SHEET]
        header = tuple(
            sheet.cell(row=4, column=c).value
            for c in range(1, len(WORKBOOK_COLUMNS) + 1)
        )
        assert header == WORKBOOK_COLUMNS


# ---------------------------------------------------------------------------
# Discovery, the read-only database browser and the export directory
# ---------------------------------------------------------------------------


class TestDiscoverEndpoint:
    def test_discovery_runs_in_the_worker_and_returns_its_report(self, client, monkeypatch):
        import regcompass.discovery as discovery_mod
        from regcompass.discovery import DiscoveryReport

        seen = {}

        def fake_discover(economy, storage, **kwargs):
            seen.update(economy=economy, refresh=kwargs.get("refresh"))
            kwargs["progress"]("Discovery | SG: spacing floor 6s/host")
            return DiscoveryReport(
                economy=economy, strategy="curl_cffi_ladder", run_id="disc_x",
                refresh=bool(kwargs.get("refresh")), fetched=2, documents_stored=2,
            )

        monkeypatch.setattr(discovery_mod, "discover_economy", fake_discover)
        r = client.post("/api/discover", json={"economy": "SG", "refresh": True})
        assert r.status_code == 200, r.text
        assert r.json()["status"] == "started"
        st = _wait_done(client)
        assert st["status"] == "done"
        assert st["mode"] == "discover"
        assert seen == {"economy": "SG", "refresh": True}
        lines = _replayed_lines(client)
        assert any("Discovery |" in line for line in lines)
        assert any("fetched 2" in line for line in lines)

    def test_a_manual_only_economy_is_refused_with_400(
        self, client, tmp_path, monkeypatch
    ):
        point_loaders_at(monkeypatch, forbidden_config(tmp_path))
        r = client.post("/api/discover", json={"economy": FORBIDDEN})
        assert r.status_code == 400
        assert "manual" in r.json()["detail"].lower()

    def test_an_unknown_economy_is_refused(self, client):
        assert client.post("/api/discover", json={"economy": "XX"}).status_code == 400

    def test_discovery_and_a_run_do_not_share_the_one_slot(self, client, monkeypatch):
        """One job at a time: a Discovery in flight refuses a second start."""
        import threading

        import regcompass.discovery as discovery_mod
        from regcompass.discovery import DiscoveryReport

        release = threading.Event()

        def slow_discover(economy, storage, **kwargs):
            release.wait(timeout=5.0)
            return DiscoveryReport(economy=economy, strategy="httpx", run_id="disc_y")

        monkeypatch.setattr(discovery_mod, "discover_economy", slow_discover)
        assert client.post("/api/discover", json={"economy": "SG"}).status_code == 200
        assert client.post("/api/discover", json={"economy": "SG"}).status_code == 409
        release.set()
        _wait_done(client)


class TestDatabaseBrowser:
    def test_it_rejects_a_non_whitelisted_table(self, client):
        assert client.get("/api/db/source_groups").status_code == 400
        assert client.get("/api/db/sqlite_master").status_code == 400
        ok = client.get("/api/db/mappings")
        assert ok.status_code == 200
        assert ok.json()["total"] == 0

    def test_pagination_bounds(self, client):
        body = client.get("/api/db/chunks?limit=100000&offset=-5").json()
        assert body["limit"] == MAX_LIMIT
        assert body["offset"] == 0
        body2 = client.get("/api/db/chunks?limit=0").json()
        assert body2["limit"] == 1
        assert "chunk_id" in body2["columns"]

    def test_verification_status_filter(self, client):
        assert client.get("/api/db/mappings?verification_status=dropped").status_code == 200
        assert client.get("/api/db/mappings?verification_status=bogus").status_code == 400
        assert client.get("/api/db/documents?verification_status=dropped").status_code == 400

    def test_a_missing_database_is_a_404(self, tmp_path):
        c = TestClient(
            create_app(
                db_path=tmp_path / "nope.db", out_dir=tmp_path,
                ui_dir=None,
            )
        )
        assert c.get("/api/db/mappings").status_code == 404


class TestOutputsDirectory:
    def test_listing_preview_and_download(self, tmp_path):
        out = tmp_path / "out"
        out.mkdir()
        (out / "submission.csv").write_text(
            "Economy,Law Name,Indicator ID\nSG,Telecommunications Act 1999,6.1\n"
            "SG,Telecommunications Act 1999,7.3\n",
            encoding="utf-8",
        )
        (out / "supplementary.json").write_text(
            json.dumps({"scores": {"SG": {"6.1": 0.5}}, "battery": "green"}),
            encoding="utf-8",
        )
        c = TestClient(
            create_app(
                db_path=tmp_path / "web.db", out_dir=out,
                ui_dir=None,
            )
        )
        listing = c.get("/api/outputs").json()
        assert {f["name"] for f in listing["files"]} == {"submission.csv", "supplementary.json"}

        csv_prev = c.get("/api/outputs/preview", params={"name": "submission.csv"}).json()
        assert csv_prev["kind"] == "csv"
        assert csv_prev["header"] == ["Economy", "Law Name", "Indicator ID"]
        assert csv_prev["n_rows_total"] == 2

        json_prev = c.get("/api/outputs/preview", params={"name": "supplementary.json"}).json()
        assert json_prev["data"]["battery"] == "green"

        dl = c.get("/api/outputs/download", params={"name": "submission.csv"})
        assert dl.status_code == 200
        assert "Economy" in dl.text

    def test_path_traversal_blocked(self, client):
        for bad in ["../pyproject.toml", "..%2Fsecret", "/etc/passwd", ".hidden"]:
            assert client.get("/api/outputs/preview", params={"name": bad}).status_code in (400, 404)

    def test_empty_when_no_dir(self, tmp_path):
        c = TestClient(
            create_app(
                db_path=tmp_path / "web.db", out_dir=tmp_path / "does_not_exist",
                ui_dir=None,
            )
        )
        assert c.get("/api/outputs").json()["files"] == []


class TestStats:
    def test_a_fresh_server_reports_zeros_and_never_500s(self, client):
        body = client.get("/api/stats").json()
        assert body["active"] is False
        assert body["status"] == "idle"
        assert body["elapsed_s"] == 0.0
        run = body["run"]
        for key in (
            "n_chunks", "n_pairs_considered", "n_pairs_gated", "n_passed",
            "n_no_evidence", "n_dropped", "n_groups", "pairs_done",
            "pairs_total", "n_documents",
        ):
            assert run[key] == 0
        assert run["economy"] is None
        assert body["rows_exported"] is None
        assert body["stages"] == []
        assert body["last_run"] is None
        assert body["models"]["names"] == []

    def test_no_flat_rate_of_any_kind(self, client):
        body = client.get("/api/stats").json()
        assert "cost" not in body and "cost_note" not in body
        assert "last_run" in body

    def test_stats_scopes_stage_times_to_this_session_s_run(self, tmp_path):
        import sqlite3 as sq
        from datetime import datetime, timedelta, timezone

        app, db, _ = _seeded_app(tmp_path)
        conn = sq.connect(db)
        conn.execute(
            "INSERT INTO audit_log (stage, method, decision, duration_ms, timestamp)"
            " VALUES ('m6_map', 'test', 'test', 999000, ?)",
            ((datetime.now(timezone.utc) - timedelta(hours=2)).isoformat(),),
        )
        conn.commit()
        conn.close()
        c = TestClient(app)
        assert c.get("/api/stats").json()["stages"] == [], "no Run yet, so no bars"
        assert c.post(
            "/api/run", json={"economy": "SG", "pillars": [7], "engine": "fake"}
        ).status_code == 200
        _wait_done(c, timeout=180.0)
        body = c.get("/api/stats").json()
        stages = {s["stage"]: s["ms"] for s in body["stages"]}
        assert stages, "this Run's stages are traced"
        assert "m6_map" not in stages or stages["m6_map"] < 999_000
        assert body["stages_scope"] == "session"
        # The Run's live meter stays on the panel once it has finished, and
        # it is the very meter the Run Record closed with.
        meter, last = body["meter"], body["last_run"]
        assert meter["calls"] > 0
        assert meter["prompt_tokens"] == last["prompt_tokens"] > 0
        assert meter["completion_tokens"] == last["completion_tokens"]
        assert meter["cost_usd"] == last["cost_usd"]
        assert body["run_id"] == last["run_id"]
        # The fake Engine runs no model, so it declares no prices to show.
        assert body["pricing"] is None

    def test_a_fresh_server_has_no_meter_and_no_prices(self, client):
        body = client.get("/api/stats").json()
        assert body["meter"] is None
        assert body["pricing"] is None
        assert body["last_run_pricing"] is None
        assert body["stages_scope"] is None
        assert body["run_id"] is None

    def test_the_meter_is_readable_while_the_run_is_still_going(self):
        """The meter lives in the worker's context; the panel reads it from a
        request thread, mid-Run, through the watcher the worker names."""
        import threading

        from regcompass.engines import record_usage, start_meter, watch_meter
        from regcompass.server import RunManager

        manager = RunManager()
        metered = threading.Event()
        release = threading.Event()

        def worker(m):
            watch_meter(m.attach_meter)
            start_meter()
            record_usage(1_000, 200)
            record_usage(500, 100, provider_cost_usd=0.001)
            metered.set()
            release.wait(5)
            m._finish("done")

        assert manager.meter_snapshot() is None
        assert manager.start(worker, {"economy": "SG", "engine": "engine-b"})
        assert metered.wait(5)
        live = manager.meter_snapshot()
        release.set()
        assert live == {
            "prompt_tokens": 1_500,
            "completion_tokens": 300,
            "provider_cost_usd": 0.001,
            "calls": 2,
        }

    def test_the_funnel_moves_during_the_run_and_counts_every_document(self):
        """The live report's counters reach /api/stats while the Run fills
        them in, and the mapped-pair bar sums every Document so far."""
        import threading

        from regcompass.pipeline import RunReport
        from regcompass.server import RunManager

        manager = RunManager()
        ready = threading.Event()
        release = threading.Event()

        def worker(m):
            report = RunReport(economy="SG", engine="fake")
            m.attach_report(report)
            report.n_chunks = 40
            report.n_pairs_considered = 400
            report.n_pairs_gated = 30
            report.n_passed = 9
            report.n_no_evidence = 3
            report.n_dropped = 1
            m.record_pairs("doc_a", 20, 20)
            m.record_pairs("doc_b", 5, 10)
            ready.set()
            release.wait(5)
            m._finish("done")

        assert manager.start(worker, {"economy": "SG", "engine": "fake"})
        assert ready.wait(5)
        live = manager.telemetry_snapshot()
        release.set()
        assert live["n_pairs_considered"] == 400
        assert live["n_pairs_gated"] == 30
        assert (live["n_passed"], live["n_no_evidence"], live["n_dropped"]) == (9, 3, 1)
        assert (live["pairs_done"], live["pairs_total"]) == (25, 30)
        assert live["n_documents"] == 2

    def test_the_meter_is_priced_at_the_engine_s_declared_rates(self):
        from regcompass.engines import cost_usd_for, resolve_engine
        from regcompass.server import _meter_block, _pricing_block

        engine = resolve_engine("engine-b")
        pricing = _pricing_block("engine-b")
        assert pricing["display_name"] == engine.display_name
        assert pricing["model"] == engine.litellm_model
        assert pricing["usd_per_million_input_tokens"] == engine.usd_per_million_input_tokens
        assert pricing["usd_per_million_output_tokens"] == engine.usd_per_million_output_tokens
        block = _meter_block(
            {"prompt_tokens": 2_000_000, "completion_tokens": 1_000_000,
             "calls": 7, "provider_cost_usd": None},
            "engine-b",
        )
        assert block["calls"] == 7
        assert block["cost_usd"] == pytest.approx(
            cost_usd_for(engine, 2_000_000, 1_000_000)
        )
        assert block["provider_cost_usd"] is None
        assert _meter_block(None, "engine-b") is None

    def test_a_restarted_server_traces_the_last_run_over_its_own_window(
        self, scratch_db, tmp_path
    ):
        """No Run in this session: the last completed Run is traced, and only
        the audit rows inside its recorded window count."""
        import sqlite3 as sq

        storage = Storage(scratch_db)
        storage.run_start(
            run_id="run_window", kind="run", economy="SG", pillars=[7],
            indicators=None, engine="engine-b",
            started_at="2026-09-01T10:00:00.000000Z",
        )
        storage.run_finish(
            "run_window", status="completed",
            ended_at="2026-09-01T10:30:00.000000Z",
            prompt_tokens=10, completion_tokens=5, cost_usd=0.01,
        )
        storage.conn.close()
        conn = sq.connect(scratch_db)
        rows = [
            ("m1_extract", 100, "2026-09-01T09:59:59.000000+00:00"),  # before
            ("m1_extract", 40, "2026-09-01T10:00:01.000000+00:00"),
            ("m6_map", 900, "2026-09-01T10:10:00.000000+00:00"),
            ("m7_verify", 0, "2026-09-01T10:10:01.000000+00:00"),
            ("m6_map", 5_000, "2026-09-01T11:00:00.000000+00:00"),  # after
        ]
        conn.executemany(
            "INSERT INTO audit_log (stage, method, decision, duration_ms, timestamp)"
            " VALUES (?, 'test', 'test', ?, ?)",
            rows,
        )
        conn.commit()
        conn.close()
        out = tmp_path / "out"
        out.mkdir()
        c = TestClient(create_app(db_path=scratch_db, out_dir=out, ui_dir=None))
        body = c.get("/api/stats").json()
        assert body["stages_scope"] == "last_run"
        assert {s["stage"]: s["ms"] for s in body["stages"]} == {
            "m1_extract": 40, "m6_map": 900, "m7_verify": 0,
        }
        assert body["meter"] is None, "no Run of this session has metered"
        assert body["last_run"]["run_id"] == "run_window"
        assert body["last_run_pricing"]["engine"] == "engine-b"

    def test_the_read_only_path_does_not_write(self, scratch_db, client):
        before = scratch_db.stat()
        assert client.get("/api/stats").status_code == 200
        after = scratch_db.stat()
        assert (before.st_size, before.st_mtime) == (after.st_size, after.st_mtime)


class TestE2ELane:
    def test_e2e_dispatches_the_e2e_worker_with_the_selected_engine(
        self, scratch_db, tmp_path, monkeypatch
    ):
        """mode=e2e with an Engine dispatches run_e2e (monkeypatched here so no
        Discovery and no models run), carrying the selected Engine object and
        the data dir the Corpus lives under."""
        import regcompass.pipeline as pipeline_mod
        from regcompass.pipeline import E2EReport, RunReport

        called = {}

        def fake_run_e2e(storage, economy, pillars, engine, data_dir, outdir, **kwargs):
            called["args"] = (economy, list(pillars), engine.name, kwargs.get("max_documents"))
            called["data_dir"] = data_dir
            kwargs["progress"]("M10 Discovery | fake e2e worker reached")
            rep = E2EReport(economy=economy, crawl_fetched=1, documents_mapped=["doc_x"])
            rep.run = RunReport(economy=economy, engine=engine.name, n_passed=1, n_groups=1)
            return rep

        monkeypatch.setattr(pipeline_mod, "run_e2e", fake_run_e2e)
        out = tmp_path / "out"
        out.mkdir()
        data = tmp_path / "crawl_data"
        c = TestClient(
            create_app(
                db_path=scratch_db, out_dir=out, data_dir=data,
                ui_dir=None,
            )
        )
        r = c.post(
            "/api/run",
            json={"economy": "SG", "pillars": [7], "engine": "fake", "mode": "e2e",
                  "max_documents": 3},
        )
        assert r.status_code == 200, r.text
        assert r.json()["corpus_empty"] is True, "e2e starts with Discovery"
        st = _wait_done(c)
        assert st["status"] == "done"
        assert called["args"] == ("SG", [7], "fake", 3)
        assert called["data_dir"] == Path(str(data))

    def test_e2e_carries_the_requested_concurrency(
        self, scratch_db, tmp_path, monkeypatch
    ):
        """The e2e lane maps through the same pool the Corpus lane does, so an
        operator asking for four calls in flight must not have it dropped on
        the way to run_e2e."""
        import regcompass.pipeline as pipeline_mod
        from regcompass.pipeline import E2EReport, RunReport

        called = {}

        def fake_run_e2e(storage, economy, pillars, engine, data_dir, outdir, **kwargs):
            called["concurrency"] = kwargs.get("concurrency")
            kwargs["progress"]("M10 Discovery | fake e2e worker reached")
            rep = E2EReport(economy=economy, crawl_fetched=1, documents_mapped=["doc_x"])
            rep.run = RunReport(economy=economy, engine=engine.name, n_passed=1)
            return rep

        monkeypatch.setattr(pipeline_mod, "run_e2e", fake_run_e2e)
        out = tmp_path / "out"
        out.mkdir()
        c = TestClient(
            create_app(
                db_path=scratch_db, out_dir=out, data_dir=tmp_path / "crawl_data",
                ui_dir=None,
            )
        )
        r = c.post(
            "/api/run",
            json={"economy": "SG", "pillars": [7], "engine": "fake", "mode": "e2e",
                  "max_documents": 1, "concurrency": 3},
        )
        assert r.status_code == 200, r.text
        assert _wait_done(c)["status"] == "done"
        assert called["concurrency"] == 3


# ---------------------------------------------------------------------------
# Box 6: the committed bundle, and the interface at phone width
# ---------------------------------------------------------------------------


class TestPhoneWidth:
    """No horizontal scroll at 400 px. The CSS assertion runs everywhere; the
    browser check runs only where a Chromium is already installed (the crawl
    extra puts one there), and is skipped rather than downloading one."""

    CSS = ROOT / "ui" / "src" / "styles.css"

    def test_the_fixed_two_pane_split_is_a_desktop_rule_only(self):
        css = self.CSS.read_text(encoding="utf-8")
        assert "@media (max-width: 720px)" in css, "no phone-width rule at all"
        head, phone = css.split("@media (max-width: 720px)", 1)
        # The 60/40 split and the fixed bar height belong to the wide layout;
        # the phone rule must override both, or the two panes stay side by side
        # on a 400 px screen.
        assert "flex: 0 0 60%" in head
        assert "flex-direction: column" in phone
        assert "width: 100%" in phone

    def test_the_served_page_does_not_scroll_sideways_at_400px(self, tmp_path):
        import socket
        import threading

        playwright = pytest.importorskip("playwright.sync_api")
        import uvicorn

        from regcompass.server import UI_DIST

        if not UI_DIST.is_dir():  # pragma: no cover - the bundle is committed
            pytest.skip("no built bundle")

        sock = socket.socket()
        sock.bind(("127.0.0.1", 0))
        port = sock.getsockname()[1]
        sock.close()

        app, _, _ = _seeded_app(tmp_path)
        app_with_ui = create_app(
            db_path=tmp_path / "seeded.db", out_dir=tmp_path / "out",
            data_dir=tmp_path / "data",
            ui_dir=UI_DIST,
        )
        config = uvicorn.Config(
            app_with_ui, host="127.0.0.1", port=port, log_level="warning"
        )
        server = uvicorn.Server(config)
        thread = threading.Thread(target=server.run, daemon=True)
        thread.start()
        try:
            deadline = time.time() + 15.0
            while not server.started and time.time() < deadline:
                time.sleep(0.05)
            assert server.started, "the server did not come up"

            with playwright.sync_playwright() as p:
                try:
                    browser = p.chromium.launch()
                except Exception as exc:  # noqa: BLE001 - no browser installed
                    pytest.skip(f"no Chromium available: {exc}")
                try:
                    page = browser.new_page(viewport={"width": 400, "height": 800})
                    page.goto(f"http://127.0.0.1:{port}/", wait_until="networkidle")
                    page.wait_for_selector("text=Run panel", timeout=10_000)
                    # At phone width the navigation is behind the Menu button.
                    for label in ("Start a Run", "Run history", "Settings"):
                        page.click("button[aria-label=Menu]")
                        page.click(f".phone-menu button:has-text('{label}')")
                        page.wait_for_timeout(200)
                        overflow = page.evaluate(
                            "() => document.documentElement.scrollWidth"
                            " - window.innerWidth"
                        )
                        assert overflow <= 0, f"{label} overflows by {overflow}px"
                finally:
                    browser.close()
        finally:
            server.should_exit = True
            thread.join(timeout=10)


# ---------------------------------------------------------------------------
# Importing the module opens nothing: the sweep belongs to a started app
# ---------------------------------------------------------------------------


class TestImportingTheModuleTouchesNoDatabase:
    """Merely importing the server module used to build an application on the
    default working database and sweep it, so a test run beside a live Run
    marked that Run interrupted while it was still going. An import must be
    inert; the sweep belongs to the moment an application actually starts.
    """

    @staticmethod
    def _db_with_a_running_run(db: Path) -> None:
        from regcompass.storage import utc_now_z

        storage = Storage(db)
        storage.apply_schema()
        storage.run_start(
            run_id="run-live-0001", kind="run", economy="MY", pillars=[6],
            indicators=None, engine="fake", started_at=utc_now_z(),
        )
        storage.close()

    @staticmethod
    def _status_of(db: Path, run_id: str = "run-live-0001") -> str:
        storage = Storage(db)
        try:
            row = storage.conn.execute(
                "SELECT status FROM runs WHERE run_id = ?", (run_id,)
            ).fetchone()
        finally:
            storage.close()
        return row["status"]

    def test_an_import_from_a_working_directory_leaves_the_run_running(
        self, tmp_path
    ):
        import subprocess
        import sys

        (tmp_path / "data").mkdir()
        db = tmp_path / "data" / "regcompass.db"
        self._db_with_a_running_run(db)

        done = subprocess.run(
            [sys.executable, "-c", "import regcompass.server"],
            cwd=tmp_path, capture_output=True, text=True, timeout=120,
        )
        assert done.returncode == 0, done.stderr
        assert self._status_of(db) == "running", done.stderr

    def test_the_module_holds_no_application_built_at_import(self):
        import regcompass.server as server_mod

        assert not hasattr(server_mod, "app"), (
            "a module-level application opens its database at import time"
        )

    def test_no_document_points_a_server_at_the_import_time_object(self):
        for name in ("README.md", "docs/WEB_DEMO.md", "docs/AUDIT_UI.md"):
            path = ROOT / name
            if not path.exists():
                continue
            assert "server:app" not in path.read_text(encoding="utf-8"), name

    def test_building_the_app_sweeps_nothing_until_it_starts(self, tmp_path):
        db = tmp_path / "server.db"
        self._db_with_a_running_run(db)
        out = tmp_path / "out"
        out.mkdir()

        app = create_app(
            db_path=db, out_dir=out, data_dir=tmp_path / "data", ui_dir=None,
        )
        assert self._status_of(db) == "running"
        assert app.state.interrupted_runs == 0

        with TestClient(app) as client:
            assert self._status_of(db) == "interrupted"
            assert client.get("/api/status").json()["interrupted_runs"] == 1
        assert app.state.interrupted_runs == 1

    def test_serve_builds_an_app_that_sweeps_when_it_starts(
        self, tmp_path, monkeypatch
    ):
        import uvicorn
        from typer.testing import CliRunner

        from regcompass.cli import app as cli_app

        db = tmp_path / "serve.db"
        self._db_with_a_running_run(db)
        out = tmp_path / "out"
        out.mkdir()

        import regcompass.server as server_mod

        real_create = server_mod.create_app
        built: dict = {}

        def spy(**kwargs):
            built["app"] = real_create(**{**kwargs, "ui_dir": None})
            return built["app"]

        monkeypatch.setattr(server_mod, "create_app", spy)
        monkeypatch.setattr(uvicorn, "run", lambda app_, **kw: None)

        r = CliRunner().invoke(
            cli_app,
            ["serve", "--db", str(db), "--out", str(out),
             "--data-dir", str(tmp_path / "data")],
        )
        assert r.exit_code == 0, r.output
        assert self._status_of(db) == "running", "the factory alone swept it"

        with TestClient(built["app"]):
            assert self._status_of(db) == "interrupted"
        assert built["app"].state.interrupted_runs == 1
