"""A corrected Mapping in the Evidence Export.

A reviewer who corrects a Mapping says: the quote is good evidence, but for
another Indicator. The export has to show that the correction took effect: the
row ships under the corrected Indicator (and so its Pillar), its Mapping
Rationale is the reviewer's reason, its Notes disclose the override, and its
Confidence stays the number the pipeline computed. The Indicator the Engine
proposed gets no evidence from that Mapping. The supplementary review_gate
record keeps the audit trail. The organizers' sheets and columns never change.

Offline, against the same hand-seeded database the review tests use (one Run
over Pillar 7, four passed Mappings 7.1 to 7.4), with an official Source URL on
the Document so the export can build its rows.
"""

from __future__ import annotations

import csv
import json
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from regcompass.export import ABSENCE_MARKER, COLUMNS, EXTRA_COLUMNS
from regcompass.storage import Storage
from regcompass.workbook import SHEET, WORKBOOK_COLUMNS, template_path

from test_reviews import DOC_ID, QUOTES, _app, _seed_document, _seed_run

REASON = "The provision sets a retention period, not a framework."


@pytest.fixture()
def client(tmp_path):
    db = tmp_path / "regcompass.db"
    storage = Storage(db)
    storage.apply_schema()
    _seed_document(storage)
    storage.upsert_document(
        DOC_ID, "SG", "sha_seeded", source_url="https://sso.agc.gov.sg/Act/SEED1999"
    )
    ids = _seed_run(storage, "run_a")
    storage.close()
    return TestClient(_app(db, tmp_path)), "run_a", ids


def _decide(c, run_id, mapping_id, status, **extra):
    r = c.post(
        "/api/reviews",
        json={"run_id": run_id, "mapping_id": mapping_id, "review_status": status, **extra},
    )
    assert r.status_code == 200, r.text


def _correct(c, run_id, mapping_id, to, reason=REASON, reviewer="ryan"):
    _decide(c, run_id, mapping_id, "corrected", corrected_indicator_id=to,
            comment=reason, reviewer=reviewer)


def _export(c, run_id):
    r = c.post("/api/export", params={"run_id": run_id})
    assert r.status_code == 200, r.text
    summary = r.json()
    with Path(summary["csv_path"]).open(encoding="utf-8-sig", newline="") as f:
        reader = csv.reader(f)
        header = next(reader)
        rows = [dict(zip(header, row)) for row in reader]
    supplementary = json.loads(Path(summary["supplementary_path"]).read_text(encoding="utf-8"))
    return summary, header, rows, supplementary


def _provisions(rows):
    return [r for r in rows if r["Article / Section"] != ABSENCE_MARKER]


def _row_for_quote(rows, quote):
    found = [r for r in _provisions(rows) if r["Verbatim Snippet"] == quote]
    assert len(found) == 1, found
    return found[0]


class TestTheCorrectedRow:
    def test_ships_under_the_corrected_indicator_with_the_reason_and_the_disclosure(self, client):
        c, run_id, ids = client
        _correct(c, run_id, ids[0], "7.5")
        _, _, rows, _ = _export(c, run_id)
        row = _row_for_quote(rows, QUOTES[0])
        assert row["Indicator ID"] == "7.5"
        assert row["Mapping Rationale"] == REASON
        assert "Reviewer override: Engine proposed 7.1; corrected to 7.5 by ryan" in row["Notes"]
        # Everything the correction does not touch stays as the Engine found it.
        assert row["Source URL"] == "https://sso.agc.gov.sg/Act/SEED1999"
        assert row["Economy"] == "Singapore"

    def test_keeps_the_pipeline_confidence(self, client):
        c, run_id, ids = client
        _decide(c, run_id, ids[0], "accepted")
        _, _, rows, _ = _export(c, run_id)
        accepted_confidence = _row_for_quote(rows, QUOTES[0])["Confidence"]
        _correct(c, run_id, ids[0], "7.5")
        _, _, rows, _ = _export(c, run_id)
        assert _row_for_quote(rows, QUOTES[0])["Confidence"] == accepted_confidence

    def test_the_disclosure_joins_the_existing_notes(self, client):
        c, run_id, ids = client
        _decide(c, run_id, ids[0], "accepted")
        _, _, rows, _ = _export(c, run_id)
        before = _row_for_quote(rows, QUOTES[0])["Notes"]
        _correct(c, run_id, ids[0], "7.5")
        _, _, rows, _ = _export(c, run_id)
        notes = _row_for_quote(rows, QUOTES[0])["Notes"]
        disclosure = "Reviewer override: Engine proposed 7.1; corrected to 7.5 by ryan"
        assert notes == (f"{before}; {disclosure}" if before else disclosure)

    def test_an_unnamed_reviewer_is_named_as_such(self, client):
        c, run_id, ids = client
        _correct(c, run_id, ids[0], "7.5", reviewer=None)
        _, _, rows, _ = _export(c, run_id)
        assert (
            "Reviewer override: Engine proposed 7.1; corrected to 7.5 by an unnamed reviewer"
            in _row_for_quote(rows, QUOTES[0])["Notes"]
        )

    def test_a_300_character_reason_ships_whole(self, client):
        c, run_id, ids = client
        reason = ("word " * 60).strip()[:299] + "."
        assert len(reason) == 300
        _correct(c, run_id, ids[0], "7.5", reason=reason)
        _, _, rows, _ = _export(c, run_id)
        assert _row_for_quote(rows, QUOTES[0])["Mapping Rationale"] == reason

    def test_joins_an_indicator_that_already_has_accepted_evidence(self, client):
        c, run_id, ids = client
        _decide(c, run_id, ids[1], "accepted")  # 7.2, its own provision
        _correct(c, run_id, ids[0], "7.2")
        _, _, rows, _ = _export(c, run_id)
        under_72 = [r for r in _provisions(rows) if r["Indicator ID"] == "7.2"]
        assert sorted(r["Verbatim Snippet"] for r in under_72) == sorted(QUOTES[:2])
        assert not [
            r for r in rows if r["Indicator ID"] == "7.2" and r["Article / Section"] == ABSENCE_MARKER
        ]


class TestTheOriginalIndicator:
    def test_gets_no_row_from_the_corrected_mapping_and_says_evidence_was_not_accepted(
        self, client
    ):
        c, run_id, ids = client
        _correct(c, run_id, ids[0], "7.5")
        _, _, rows, _ = _export(c, run_id)
        under_71 = [r for r in rows if r["Indicator ID"] == "7.1"]
        assert len(under_71) == 1
        assert under_71[0]["Article / Section"] == ABSENCE_MARKER
        assert "1 verified Mapping for this indicator was found but not accepted in review" in under_71[0]["Notes"]


class TestTheReviewGateRecord:
    def test_lists_every_override_and_the_counts_add_up(self, client):
        c, run_id, ids = client
        _correct(c, run_id, ids[0], "7.5")
        _decide(c, run_id, ids[1], "accepted")
        _decide(c, run_id, ids[2], "rejected")
        summary, _, _, supplementary = _export(c, run_id)
        gate = supplementary["review_gate"]
        assert gate["n_corrected"] == 1
        assert (gate["n_accepted"], gate["n_rejected"], gate["n_flagged"]) == (1, 1, 0)
        assert gate["n_unreviewed"] == 1
        assert (
            gate["n_accepted"] + gate["n_corrected"] + gate["n_rejected"]
            + gate["n_flagged"] + gate["n_unreviewed"]
        ) == gate["n_verified"] == 4
        assert "corrected" in gate["rule"]
        [override] = gate["overrides"]
        decided_at = c.get("/api/reviews", params={"run_id": run_id}).json()["reviews"]
        decided_at = next(r["reviewed_at"] for r in decided_at if r["mapping_id"] == ids[0])
        assert override == {
            "mapping_id": ids[0],
            "original_indicator_id": "7.1",
            "corrected_indicator_id": "7.5",
            "reviewer": "ryan",
            "decided_at": decided_at,
            "reason": REASON,
        }
        assert summary["n_corrected"] == 1

    def test_no_correction_means_an_empty_override_list(self, client):
        c, run_id, ids = client
        _decide(c, run_id, ids[0], "accepted")
        _, _, _, supplementary = _export(c, run_id)
        gate = supplementary["review_gate"]
        assert gate["n_corrected"] == 0
        assert gate["overrides"] == []


class TestTheTemplateIsUnchanged:
    def test_csv_columns_workbook_sheets_and_headers_are_the_organizers(self, client):
        c, run_id, ids = client
        _correct(c, run_id, ids[0], "7.5")
        summary, header, _, _ = _export(c, run_id)
        assert header == list(COLUMNS) + list(EXTRA_COLUMNS)

        template = openpyxl.load_workbook(template_path())
        written = openpyxl.load_workbook(summary["xlsx_path"])
        assert written.sheetnames == template.sheetnames
        ws, tws = written[SHEET], template[SHEET]
        width = len(WORKBOOK_COLUMNS) + 1  # A..N plus their Pillar formula in O
        for row in range(1, 6):
            assert [ws.cell(row=row, column=col).value for col in range(1, width + 1)] == [
                tws.cell(row=row, column=col).value for col in range(1, width + 1)
            ]
        # The corrected row sits under 7.5, and the Pillar formula reads it.
        entry = [
            r for r in range(9, 110)
            if ws.cell(row=r, column=WORKBOOK_COLUMNS.index("Verbatim Snippet") + 1).value
            == QUOTES[0]
        ]
        assert len(entry) == 1
        assert ws.cell(row=entry[0], column=WORKBOOK_COLUMNS.index("Indicator ID") + 1).value == "7.5"
        assert ws.cell(row=entry[0], column=width).value == tws.cell(row=entry[0], column=width).value


class TestTheExportPreview:
    def test_counts_corrected_on_its_own(self, client):
        c, run_id, ids = client
        _correct(c, run_id, ids[0], "7.5")
        _decide(c, run_id, ids[1], "accepted")
        preview = c.get("/api/export/preview", params={"run_id": run_id}).json()
        assert (preview["n_accepted"], preview["n_corrected"], preview["n_unreviewed"]) == (1, 1, 2)
