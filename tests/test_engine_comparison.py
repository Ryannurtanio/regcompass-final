"""The organizers' Engine Comparison sheet: every provision either Engine
cited, paired across the two Engines, plus a per-Engine summary of each pass.

The sheet (OUTPUT_TEMPLATE_FINAL_ROUND.xlsx, sheet 4) is handed in after the
live hour. Everything here is offline: seeded Run Records and Mappings, no
model, and the organizers' template read from the vendored copy, never
written."""

from __future__ import annotations

import csv
import hashlib
import io

from openpyxl import load_workbook

from regcompass.compare import (
    PROVISION_MATCH_BASIS,
    citation_key,
    compare_provisions,
    compare_runs,
    engine_sheet_csv,
    engine_summary,
    hour_discoveries,
    pass_discoveries,
    provision_number,
)
from regcompass.contracts import RunRecord
from regcompass.workbook import (
    ENGINE_CONTINUATION_SHEET,
    ENGINE_PROVISION_HEADERS,
    ENGINE_SHEET,
    ENGINE_SUMMARY_FIELDS,
    template_path,
    write_engine_comparison,
)

from test_compare import DOC_A, DOC_B, TITLES, a_mapping, a_run, paired_db  # noqa: F401

QUOTE = "No organisation shall process personal data without consent."


def rows_of(mappings_a, mappings_b, indicators=("7.1", "7.2", "7.3", "7.4", "7.5")):
    return compare_provisions(mappings_a, mappings_b, indicators, document_titles=TITLES)


def a_discovery(run_id, *, started_at, ended_at, fetched, economy="SG", **kw):
    return RunRecord(
        run_id=run_id, kind="discovery", economy=economy, pillars=[],
        status="completed", started_at=started_at, ended_at=ended_at,
        documents_fetched=fetched, **kw,
    )


# ---------------------------------------------------------------------------
# how a provision is matched
# ---------------------------------------------------------------------------


class TestProvisionIdentity:
    def test_the_number_is_read_whatever_the_label_spelling(self):
        for label in ("s. 26", "S.26", "Section 26", "section 26(1)", "Art. 26",
                      "Article 26", "Pasal 26", "26"):
            assert provision_number(label) == "26", label
        assert provision_number("s. 26A") == "26a"
        assert provision_number("第二十六条") == "二十六"
        assert provision_number("Schedule  Two") == "schedule two"

    def test_a_citation_ignores_the_prefix_and_spacing_but_not_the_subdivision(self):
        one = a_mapping("7.1", DOC_A, "Section 26", subsection="(1)")
        two = a_mapping("7.1", DOC_A, "s. 26", subsection=" (1)")
        whole = a_mapping("7.1", DOC_A, "s. 26")
        other = a_mapping("7.1", DOC_A, "s. 26", subsection="(2)")
        assert citation_key(one) == citation_key(two)
        assert citation_key(one) != citation_key(whole)
        assert citation_key(one) != citation_key(other)

    def test_the_basis_is_named(self):
        assert "s. 26(1)" in PROVISION_MATCH_BASIS


class TestProvisionRows:
    def test_the_organizers_example_row(self):
        # Engine A cited s. 26(1), Engine B cited s. 26, same text
        [row] = rows_of(
            [a_mapping("7.3", DOC_A, "s. 26", subsection="(1)", quote=QUOTE)],
            [a_mapping("7.3", DOC_A, "s. 26", quote=QUOTE)],
        )
        assert (row.found_by, row.indicator_differs, row.citation_differs,
                row.quoted_words_differ) == ("Both", "No", "Yes", "No")
        assert row.article_section == "s. 26"
        assert row.indicator_id == "7.3"
        assert row.law_name == TITLES[DOC_A]
        assert row.difference == (
            "Engine A cited s. 26(1), Engine B cited s. 26,"
            " A is the more precise citation; same quoted words"
        )

    def test_a_shared_article_is_named_in_the_drafting_word(self):
        """Same Article, different subdivisions: the row names the Article the
        way the Economy drafts it, as the Difference line beside it does."""
        [row] = rows_of(
            [a_mapping("7.1", DOC_A, "Chapter 2 s. 7", subsection="(4)").model_copy(
                update={"economy": "CN"})],
            [a_mapping("7.1", DOC_A, "Chapter 2 s. 7", subsection="(1)").model_copy(
                update={"economy": "CN"})],
        )
        assert row.citation_differs == "Yes"
        assert row.article_section == "Chapter 2 Art. 7"

    def test_full_agreement_says_so(self):
        [row] = rows_of(
            [a_mapping("7.1", DOC_A, "s. 13", subsection="(1)")],
            [a_mapping("7.1", DOC_A, "S.  13", subsection="(1)")],
        )
        assert (row.indicator_differs, row.citation_differs, row.quoted_words_differ) == (
            "No", "No", "No")
        assert row.difference == "Same provision, Indicator and quoted words on both Engines"

    def test_quoted_words_ignore_case_spacing_and_punctuation(self):
        [row] = rows_of(
            [a_mapping("7.1", DOC_A, "s. 13", quote="No organisation shall  process data.")],
            [a_mapping("7.1", DOC_A, "s. 13", quote="no organisation shall process data")],
        )
        assert row.quoted_words_differ == "No"

    def test_one_quote_inside_the_other_is_named(self):
        [row] = rows_of(
            [a_mapping("7.1", DOC_A, "s. 13", quote="(1) No organisation shall process data.")],
            [a_mapping("7.1", DOC_A, "s. 13", quote="No organisation shall process data")],
        )
        assert row.quoted_words_differ == "Yes"
        assert row.difference == "Engine A quoted more of the same text"

    def test_a_different_indicator_on_one_provision_is_one_row(self):
        [row] = rows_of(
            [a_mapping("7.1", DOC_A, "s. 13")],
            [a_mapping("7.2", DOC_A, "s. 13")],
        )
        assert row.found_by == "Both"
        assert row.indicator_differs == "Yes"
        assert row.indicator_id == "7.1 / 7.2"
        assert "Engine A mapped it to 7.1, Engine B to 7.2" in row.difference

    def test_the_same_indicator_is_paired_before_a_different_one(self):
        rows = rows_of(
            [a_mapping("7.1", DOC_A, "s. 13"), a_mapping("7.2", DOC_A, "s. 13")],
            [a_mapping("7.2", DOC_A, "s. 13")],
        )
        by = {(r.found_by, r.indicator_id): r for r in rows}
        assert set(by) == {("Engine A only", "7.1"), ("Both", "7.2")}
        assert by[("Both", "7.2")].indicator_differs == "No"
        assert by[("Engine A only", "7.1")].indicator_differs == "n/a"
        assert by[("Engine A only", "7.1")].citation_differs == "n/a"
        assert by[("Engine A only", "7.1")].quoted_words_differ == "n/a"

    def test_a_provision_one_engine_cited_names_what_the_other_chose(self):
        rows = rows_of(
            [a_mapping("7.1", DOC_A, "s. 13"), a_mapping("7.3", DOC_A, "s. 24")],
            [a_mapping("7.1", DOC_B, "s. 5", subsection="(2)")],
        )
        by = {(r.found_by, r.indicator_id, r.article_section): r for r in rows}
        assert by[("Engine A only", "7.1", "s. 13")].difference == (
            "Only Engine A cited this provision; for 7.1 Engine B cited s. 5(2)"
            " of Cybersecurity Act 2018"
        )
        assert by[("Engine B only", "7.1", "s. 5(2)")].difference == (
            "Only Engine B cited this provision; for 7.1 Engine A cited s. 13"
            " of Personal Data Protection Act 2012"
        )
        assert by[("Engine A only", "7.3", "s. 24")].difference == (
            "Only Engine A cited this provision; Engine B found no evidence for 7.3"
        )

    def test_every_evidence_mapping_is_a_row_not_only_the_controlling_one(self):
        rows = rows_of(
            [a_mapping("7.1", DOC_A, "s. 13", controlling=True),
             a_mapping("7.1", DOC_A, "s. 14", controlling=False, chunk_id="c14")],
            [],
        )
        assert [r.article_section for r in rows] == ["s. 13", "s. 14"]
        assert {r.found_by for r in rows} == {"Engine A only"}

    def test_dropped_and_no_evidence_mappings_are_not_provisions(self):
        rows = rows_of(
            [a_mapping("7.1", DOC_A, "s. 13", verification_status="dropped")],
            [],
        )
        assert rows == []

    def test_one_passage_quoted_from_two_chunks_is_one_provision(self):
        rows = rows_of(
            [a_mapping("7.1", DOC_A, "s. 13", chunk_id="c1"),
             a_mapping("7.1", DOC_A, "s. 13", chunk_id="c2")],
            [a_mapping("7.1", DOC_A, "s. 13")],
        )
        assert [r.found_by for r in rows] == ["Both"]

    def test_indicators_outside_the_list_are_not_read(self):
        assert rows_of([a_mapping("6.1", DOC_A, "s. 13")], [], indicators=("7.1",)) == []

    def test_rows_are_in_indicator_then_section_order(self):
        rows = rows_of(
            [a_mapping("7.2", DOC_A, "s. 9"), a_mapping("7.1", DOC_A, "s. 13"),
             a_mapping("7.1", DOC_A, "s. 2", chunk_id="c2")],
            [],
        )
        assert [(r.indicator_id, r.article_section) for r in rows] == [
            ("7.1", "s. 2"), ("7.1", "s. 13"), ("7.2", "s. 9")]

    def test_the_comparison_carries_the_rows_and_the_counts(self):
        comparison = compare_runs(
            a_run("run_a", "engine-a", started_at="2026-10-15T03:00:00Z",
                  ended_at="2026-10-15T03:20:00Z"),
            a_run("run_b", "engine-b", started_at="2026-10-15T03:30:00Z",
                  ended_at="2026-10-15T03:40:00Z"),
            [a_mapping("7.1", DOC_A, "s. 13"), a_mapping("7.2", DOC_A, "s. 9")],
            [a_mapping("7.1", DOC_A, "s. 13"), a_mapping("7.4", DOC_B, "s. 8")],
            ("7.1", "7.2", "7.3", "7.4", "7.5"),
            document_titles=TITLES,
        )
        assert comparison.n_provisions == 3
        assert (comparison.n_found_by_a_only, comparison.n_found_by_b_only,
                comparison.n_found_by_both) == (1, 1, 1)
        # the Indicator view is unchanged beside it
        assert [r.agreement for r in comparison.rows] == [
            "agree", "only_a", "neither", "only_b", "neither"]


# ---------------------------------------------------------------------------
# block 1: the per-engine summary, from the Run Records alone
# ---------------------------------------------------------------------------


RUN_A = a_run("run_a", "engine-a", started_at="2026-10-15T03:10:00Z",
              ended_at="2026-10-15T03:40:00Z", cost_usd=0.4, provider_cost_usd=0.61)
RUN_B = a_run("run_b", "engine-b", started_at="2026-10-15T03:45:00Z",
              ended_at="2026-10-15T03:55:30Z", cost_usd=0.05)
DISCOVERY = a_discovery("disc_1", started_at="2026-10-15T03:00:00Z",
                        ended_at="2026-10-15T03:08:00Z", fetched=5)
YESTERDAY_RUN = a_run("run_old", "engine-b", started_at="2026-10-14T02:00:00Z",
                      ended_at="2026-10-14T02:30:00Z")
OLD_DISCOVERY = a_discovery("disc_old", started_at="2026-10-14T01:00:00Z",
                            ended_at="2026-10-14T01:30:00Z", fetched=9)
RECORDS = [OLD_DISCOVERY, YESTERDAY_RUN, DISCOVERY, RUN_A, RUN_B]


class TestEngineSummary:
    def test_engine_a_s_pass_starts_with_its_discovery(self):
        summary = engine_summary(RUN_A, RECORDS, provider_model="openrouter / openai/gpt-5.6-luna")
        assert summary.discovery_run_ids == ["disc_1"]
        assert summary.documents_fetched == 5
        assert summary.start_hhmm == "10:00"  # Bangkok, UTC+7
        assert summary.end_hhmm == "10:40"
        assert summary.elapsed_minutes == 40.0
        assert summary.provider_model == "openrouter / openai/gpt-5.6-luna"

    def test_engine_b_fetching_nothing_shows_zero(self):
        summary = engine_summary(RUN_B, RECORDS)
        assert summary.discovery_run_ids == []
        assert summary.documents_fetched == 0
        assert summary.elapsed_minutes == 10.5

    def test_a_fetch_between_the_two_runs_is_engine_a_s(self):
        """Runs never fetch: an add made between the two Runs was fetched for
        the first pass, and the second pass still fetched nothing."""
        stray = a_discovery("disc_2", started_at="2026-10-15T03:41:00Z",
                            ended_at="2026-10-15T03:42:00Z", fetched=2)
        records = [*RECORDS, stray]
        hour = hour_discoveries(RUN_A, RUN_B, records)
        assert [r.run_id for r in hour] == ["disc_1", "disc_2"]
        a = engine_summary(RUN_A, records, discoveries=hour)
        assert a.documents_fetched == 7
        assert (a.start_hhmm, a.end_hhmm) == ("10:00", "10:42")
        assert engine_summary(RUN_B, records, second_pass=True).documents_fetched == 0
        comparison = compare_runs(RUN_A, RUN_B, [], [], ("7.3",), economy_records=records)
        assert (comparison.pass_a.documents_fetched, comparison.pass_b.documents_fetched) == (7, 0)

    def test_a_stopped_run_never_closes_the_window(self):
        """Discovery, Engine A Run stopped, Engine A restarted, then Engine B:
        the Discovery is still the first pass's."""
        stopped = a_run("run_a0", "engine-a", started_at="2026-10-15T03:09:00Z",
                        ended_at="2026-10-15T03:09:30Z", status="failed")
        records = [OLD_DISCOVERY, YESTERDAY_RUN, DISCOVERY, stopped, RUN_A, RUN_B]
        assert [r.run_id for r in pass_discoveries(RUN_A, records)] == ["disc_1"]
        assert engine_summary(
            RUN_A, records, discoveries=hour_discoveries(RUN_A, RUN_B, records)
        ).documents_fetched == 5
        # nor does a completed earlier Run on the same Engine (a re-run)
        rerun = a_run("run_a1", "engine-a", started_at="2026-10-15T03:08:30Z",
                      ended_at="2026-10-15T03:09:00Z")
        assert [r.run_id for r in pass_discoveries(RUN_A, [*records, rerun])] == ["disc_1"]

    def test_runs_started_together_claim_the_discovery_once(self):
        """Engine B started while Engine A was still running: the Discovery is
        the first pass's, and the second pass fetched nothing."""
        b_early = a_run("run_b", "engine-b", started_at="2026-10-15T03:20:00Z",
                        ended_at="2026-10-15T03:50:00Z")
        records = [DISCOVERY, RUN_A, b_early]
        comparison = compare_runs(RUN_A, b_early, [], [], ("7.3",), economy_records=records)
        assert comparison.pass_a.documents_fetched == 5
        assert comparison.pass_b.documents_fetched == 0

    def test_a_discovery_before_an_earlier_run_is_not_this_pass(self):
        assert [r.run_id for r in pass_discoveries(RUN_A, RECORDS)] == ["disc_1"]
        assert pass_discoveries(YESTERDAY_RUN, RECORDS) == [OLD_DISCOVERY]

    def test_another_economy_s_discovery_is_never_counted(self):
        elsewhere = a_discovery("disc_my", started_at="2026-10-15T03:00:00Z",
                                ended_at="2026-10-15T03:05:00Z", fetched=3, economy="MY")
        assert engine_summary(RUN_B, [*RECORDS, elsewhere]).documents_fetched == 0
        assert engine_summary(RUN_A, [*RECORDS, elsewhere]).documents_fetched == 5

    def test_cost_is_the_provider_figure_when_reported(self):
        a = engine_summary(RUN_A, RECORDS)
        b = engine_summary(RUN_B, RECORDS)
        assert (a.cost_usd, a.cost_basis) == (0.61, "provider")
        assert (b.cost_usd, b.cost_basis) == (0.05, "metered")

    def test_without_records_the_pass_is_the_run(self):
        summary = engine_summary(RUN_A)
        assert summary.documents_fetched == 0
        assert summary.elapsed_minutes == 30.0


# ---------------------------------------------------------------------------
# the organizers' sheet as an xlsx and as a csv
# ---------------------------------------------------------------------------


def a_comparison(n_extra: int = 0):
    mappings_a = [a_mapping("7.1", DOC_A, "s. 26", subsection="(1)", quote=QUOTE),
                  a_mapping("7.2", DOC_A, "s. 9")]
    mappings_b = [a_mapping("7.1", DOC_A, "s. 26", quote=QUOTE),
                  a_mapping("7.4", DOC_B, "s. 12")]
    mappings_a += [
        a_mapping("7.3", DOC_A, f"s. {100 + i}", chunk_id=f"x{i}") for i in range(n_extra)
    ]
    return compare_runs(
        RUN_A, RUN_B, mappings_a, mappings_b, ("7.1", "7.2", "7.3", "7.4", "7.5"),
        document_titles=TITLES, economy_records=RECORDS,
        engine_models={"engine-a": "openrouter / openai/gpt-5.6-luna",
                       "engine-b": "openrouter / qwen/qwen3-30b-a3b-instruct-2507"},
    )


def sheet_of(tmp_path, comparison, template=None):
    out = tmp_path / "engine.xlsx"
    result = write_engine_comparison(
        template_path() if template is None else template, out, comparison
    )
    return result, load_workbook(out)


class TestEngineComparisonWorkbook:
    def test_the_organizers_template_is_never_written(self, tmp_path):
        before = hashlib.sha256(template_path().read_bytes()).hexdigest()
        sheet_of(tmp_path, a_comparison())
        assert hashlib.sha256(template_path().read_bytes()).hexdigest() == before

    def test_the_file_is_the_organizers_sheet_filled(self, tmp_path):
        result, wb = sheet_of(tmp_path, a_comparison())
        assert result.from_template and result.n_rows == 3 and result.n_continued == 0
        assert wb.sheetnames == [ENGINE_SHEET]
        ws = wb[ENGINE_SHEET]
        assert [ws.cell(row=5 + i, column=1).value for i in range(6)] == list(
            ENGINE_SUMMARY_FIELDS)
        assert [ws.cell(row=15, column=c).value for c in range(1, 10)] == list(
            ENGINE_PROVISION_HEADERS)
        assert [ws[f"B{r}"].value for r in range(5, 11)] == [
            "openrouter / openai/gpt-5.6-luna", "10:00", "10:40", 40.0, 5, 0.61]
        assert [ws[f"C{r}"].value for r in range(5, 11)] == [
            "openrouter / qwen/qwen3-30b-a3b-instruct-2507", "10:45", "10:55", 10.5, 0, 0.05]

    def test_block_two_holds_every_provision_from_row_17(self, tmp_path):
        comparison = a_comparison()
        _, wb = sheet_of(tmp_path, comparison)
        ws = wb[ENGINE_SHEET]
        assert all(ws.cell(row=16, column=c).value is None for c in range(1, 10))
        got = [[ws.cell(row=r, column=c).value for c in range(1, 10)] for r in (17, 18, 19)]
        assert got[0] == [1, TITLES[DOC_A], "s. 26", "7.1", "Both", "No", "Yes", "No",
                          comparison.provisions[0].difference]
        assert [g[4] for g in got] == ["Both", "Engine A only", "Engine B only"]
        assert ws["D17"].number_format == "@"
        assert ws["A20"].value is None

    def test_dropdowns_counters_and_block_three_survive(self, tmp_path):
        _, wb = sheet_of(tmp_path, a_comparison())
        ws = wb[ENGINE_SHEET]
        rules = {str(dv.sqref): dv.formula1 for dv in ws.data_validations.dataValidation}
        assert rules == {"E17:E56": '"Engine A only,Engine B only,Both"',
                         "F17:H56": '"Yes,No,n/a"'}
        assert ws["D58"].value == '=COUNTIF(E17:E56,"Engine A only")'
        assert ws["D60"].value == '=COUNTIF(E17:E56,"Both")'
        assert ws["A63"].value is None
        assert "A63:I67" in {str(m) for m in ws.merged_cells.ranges}

    def test_more_than_forty_provisions_continue_on_a_second_sheet(self, tmp_path):
        comparison = a_comparison(n_extra=42)  # 45 rows
        result, wb = sheet_of(tmp_path, comparison)
        assert (result.n_rows, result.n_on_sheet, result.n_continued) == (45, 40, 5)
        assert wb.sheetnames == [ENGINE_SHEET, ENGINE_CONTINUATION_SHEET]
        ws, extra = wb[ENGINE_SHEET], wb[ENGINE_CONTINUATION_SHEET]
        assert ws["A56"].value == 40
        assert [extra.cell(row=3, column=c).value for c in range(1, 10)] == list(
            ENGINE_PROVISION_HEADERS)
        assert [extra.cell(row=r, column=1).value for r in range(4, 9)] == [41, 42, 43, 44, 45]
        assert extra["D4"].number_format == "@"
        assert ws["D58"].value == (
            '=COUNTIF(E17:E56,"Engine A only")'
            "+COUNTIF('Engine Comparison (cont.)'!E4:E8,\"Engine A only\")"
        )
        assert "continue" in result.note and "45 provisions" in ws["A57"].value
        rules = {str(dv.sqref) for dv in extra.data_validations.dataValidation}
        assert rules == {"E4:E8", "F4:H8"}

    def test_without_the_template_the_same_layout_is_built(self, tmp_path):
        result, wb = sheet_of(tmp_path, a_comparison(), template=tmp_path / "missing.xlsx")
        assert not result.from_template
        ws = wb[ENGINE_SHEET]
        assert ws["A1"].value.startswith("Finale morning")
        assert [ws.cell(row=15, column=c).value for c in range(1, 10)] == list(
            ENGINE_PROVISION_HEADERS)
        assert ws["E17"].value == "Both" and ws["B5"].value.startswith("openrouter")
        assert ws["D59"].value == '=COUNTIF(E17:E56,"Engine B only")'
        rules = {str(dv.sqref) for dv in ws.data_validations.dataValidation}
        assert rules == {"E17:E56", "F17:H56"}

    def test_the_csv_is_the_sheet_in_its_own_words(self):
        comparison = a_comparison()
        rows = list(csv.reader(io.StringIO(engine_sheet_csv(comparison))))
        assert rows[0] == ["Field", "Engine A — first pass", "Engine B — second pass"]
        assert rows[5] == ["Documents fetched during this pass", "5", "0"]
        header = rows.index(list(ENGINE_PROVISION_HEADERS))
        assert rows[header + 1][:8] == [
            "1", TITLES[DOC_A], "s. 26", "7.1", "Both", "No", "Yes", "No"]
        assert ["Found by both:", "1"] in rows


# ---------------------------------------------------------------------------
# the server: payload and downloads
# ---------------------------------------------------------------------------


class TestEngineComparisonApi:
    def test_the_payload_carries_the_passes_and_the_provisions(self, paired_db):
        client, _ = paired_db
        payload = client.get(
            "/api/compare", params={"run_a": "run_a", "run_b": "run_b"}).json()
        assert payload["pass_a"]["provider_model"] == "openrouter / openai/gpt-5.6-luna"
        assert payload["pass_b"]["provider_model"] == (
            "openrouter / qwen/qwen3-30b-a3b-instruct-2507")
        assert payload["pass_b"]["documents_fetched"] == 0
        # 7.1 and 7.2 both, 7.2 on another subsection; 7.3 A only; 7.4 B only
        assert payload["n_provisions"] == len(payload["provisions"]) == 4
        assert {p["found_by"] for p in payload["provisions"]} == {
            "Both", "Engine A only", "Engine B only"}

    def test_the_xlsx_download_is_the_filled_sheet(self, paired_db):
        client, _ = paired_db
        r = client.get("/api/compare/download",
                       params={"run_a": "run_a", "run_b": "run_b", "format": "xlsx"})
        assert r.status_code == 200
        assert r.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        assert "comparison_SG_p7_engine-a_vs_engine-b.engine-comparison.xlsx" in (
            r.headers["content-disposition"])
        ws = load_workbook(io.BytesIO(r.content))[ENGINE_SHEET]
        assert ws["B5"].value == "openrouter / openai/gpt-5.6-luna"
        assert ws["C9"].value == 0
        assert ws["E17"].value in ("Both", "Engine A only", "Engine B only")

    def test_the_sheet_csv_download(self, paired_db):
        client, _ = paired_db
        r = client.get("/api/compare/download",
                       params={"run_a": "run_a", "run_b": "run_b", "format": "sheet_csv"})
        assert r.status_code == 200
        assert r.headers["content-type"].startswith("text/csv")
        assert "engine-comparison.csv" in r.headers["content-disposition"]
        assert "Found by Engine A only:" in r.text

    def test_a_discovery_before_run_a_is_counted_on_engine_a(self, paired_db):
        from regcompass.storage import Storage

        client, path = paired_db
        storage = Storage(path)
        storage.run_start(run_id="disc_1", kind="discovery", economy="SG", pillars=[],
                          indicators=None, engine=None,
                          started_at="2026-08-31T23:50:00Z")
        storage.run_finish("disc_1", status="completed", ended_at="2026-08-31T23:55:00Z",
                           documents_fetched=4)
        storage.conn.close()
        payload = client.get(
            "/api/compare", params={"run_a": "run_a", "run_b": "run_b"}).json()
        assert payload["pass_a"]["documents_fetched"] == 4
        assert payload["pass_a"]["discovery_run_ids"] == ["disc_1"]
        assert payload["pass_b"]["documents_fetched"] == 0


class TestFrameworkNote:
    """7.1 and 7.2 reach the evidence file once per Economy and only on a law
    of the right family, so a side with no such law ships no row for it; the
    Engine Comparison says so where the two Engines are set side by side."""

    def test_a_side_without_a_framework_law_is_named(self):
        # Engine A tagged 7.2 on the data-protection Act; Engine B has no 7.2
        comparison = a_comparison()
        side = (
            " has no 7.2 Mapping on a cybersecurity law or a law the 2025 baseline"
            " cites for 7.2, so its evidence file has no 7.2 row"
        )
        assert comparison.framework_note == f"Engine A{side}; Engine B{side}."

    def test_both_sides_on_the_right_laws_say_nothing(self):
        mappings = [a_mapping("7.1", DOC_A, "s. 26"), a_mapping("7.2", DOC_B, "s. 3")]
        comparison = compare_runs(
            RUN_A, RUN_B, mappings, list(mappings), ("7.1", "7.2"),
            document_titles=TITLES, economy_records=RECORDS,
        )
        assert comparison.framework_note is None

    def test_a_pillar_without_7_1_or_7_2_says_nothing(self):
        mappings = [a_mapping("7.4", DOC_B, "s. 12")]
        comparison = compare_runs(
            RUN_A, RUN_B, mappings, list(mappings), ("7.3", "7.4"),
            document_titles=TITLES, economy_records=RECORDS,
        )
        assert comparison.framework_note is None

    def test_the_note_is_written_in_row_57(self, tmp_path):
        comparison = a_comparison()
        _, wb = sheet_of(tmp_path, comparison)
        assert wb[ENGINE_SHEET]["A57"].value == comparison.framework_note
