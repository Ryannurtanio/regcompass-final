"""The Evidence Export as the organizer's own workbook (M9 / final round).

Lanes: the vendored template is a byte-identical copy of the organizers' file;
the writer deletes the two example rows as the Instructions ask, shifting the
Pillar formula, the four validations, the autofilter and the Coverage Matrix
bounds up with them exactly as a spreadsheet application would, then fills
Output Data from row 7 and never past row 107; Indicator IDs are TEXT cells (4.01 and
12.4.1 must survive); Language of Source is one of the organizer's eleven
values, with "Other" plus a Notes disclosure when nothing honest is known; over
the 101-row cap, accepted rows are picked round-robin over Economies by
Confidence so every Economy appears, and the cut is reported.
"""

from __future__ import annotations

import csv
import hashlib
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook

from regcompass.config import load_portals
from regcompass.contracts import (
    ORGANIZER_LANGUAGES,
    CorpusDoc,
    MappingRecord,
    organizer_language,
)
from regcompass.export import (
    COLUMNS,
    EXTRA_COLUMNS,
    LANGUAGE_COLUMN,
    LANGUAGE_UNKNOWN_NOTE,
    SyntheticDoc,
    export_all,
    select_rows,
)
from regcompass.workbook import (
    EXAMPLE_ROWS,
    FIRST_ROW,
    LAST_ROW,
    ROW_CAP,
    SHEET,
    WORKBOOK_COLUMNS,
    template_path,
    write_workbook,
)

ROOT = Path(__file__).resolve().parents[1]
ORGANIZER_DIR = ROOT / "config/organizer"


# ---------------------------------------------------------------------------
# the vendored template
# ---------------------------------------------------------------------------


class TestVendoredTemplate:
    def test_sha256_matches_the_recorded_value(self):
        """config/organizer/README.md records the SHA-256 of the organizers'
        file. The copy must still be byte-identical: a template that drifted
        silently would ship a workbook the secretariat never wrote."""
        readme = (ORGANIZER_DIR / "README.md").read_text(encoding="utf-8")
        m = re.search(r"\b([0-9a-f]{64})\b", readme)
        assert m, "README.md must record the template SHA-256"
        actual = hashlib.sha256(template_path().read_bytes()).hexdigest()
        assert actual == m.group(1)

    def test_template_path_resolves_under_the_config_dir(self, tmp_path):
        assert template_path().parent == ORGANIZER_DIR
        assert template_path(tmp_path) == tmp_path / "organizer" / template_path().name

    def test_workbook_columns_are_the_csv_contract_plus_language(self):
        assert WORKBOOK_COLUMNS == COLUMNS + (LANGUAGE_COLUMN,)
        assert len(WORKBOOK_COLUMNS) == 14  # A..N; O is the organizers' formula

    def test_workbook_columns_are_the_templates_own_headers(self):
        """The organizers' header row is the contract. If their A..N ever stop
        matching our column order, the writer must fail here, not in a file the
        secretariat validates."""
        ws = load_workbook(template_path())[SHEET]
        headers = tuple(
            ws.cell(row=4, column=c).value for c in range(1, len(WORKBOOK_COLUMNS) + 1)
        )
        assert headers == WORKBOOK_COLUMNS


# ---------------------------------------------------------------------------
# the writer
# ---------------------------------------------------------------------------


def wb_row(**over) -> dict:
    row = {
        "Economy": "Indonesia",
        "Law Name": "Personal Data Protection Law 2022",
        "Law Number / Ref": "Law No. 27/2022",
        "Last Amended": "2022",
        "Indicator ID": "6.1",
        "Article / Section": "Art. 56(1)",
        "Discovery Tag": "NEW",
        "Location Reference": "PDF: page 12",
        "Verbatim Snippet": "Personal data must be processed within the territory.",
        "Mapping Rationale": "This Art. 56(1) requires local processing.",
        "Source URL": "https://peraturan.go.id/id/uu-no-27-tahun-2022",
        "Confidence": "0.81",
        "Notes": "",
        LANGUAGE_COLUMN: "Bahasa Indonesia",
    }
    row.update(over)
    return row


def written(tmp_path, rows, **kw):
    out = tmp_path / "submission.xlsx"
    result = write_workbook(template_path(), out, rows, **kw)
    return result, load_workbook(out)


def pillar_from_formula(value: str) -> int | str:
    """The organizers' O-column formula, evaluated in Python: INT(E) when E is
    numeric, else the text before the first dot, else '?'. openpyxl stores
    formulas, it does not evaluate them, so the test computes the same thing."""
    if value == "":
        return ""
    try:
        return int(float(value))
    except ValueError:
        pass
    head = value.split(".", 1)[0]
    try:
        return int(head)
    except ValueError:
        return "?"


class TestWriteWorkbook:
    def test_rows_land_from_row_7_in_the_organizers_column_order(self, tmp_path):
        rows = [wb_row(), wb_row(**{"Indicator ID": "7.3", "Article / Section": "Art. 60"})]
        result, wb = written(tmp_path, rows)
        ws = wb[SHEET]
        assert result.n_rows == 2 and result.rows_cut == 0
        assert FIRST_ROW == 7
        assert ws["A7"].value == "Indonesia"
        assert ws["F8"].value == "Art. 60"
        assert ws["N7"].value == "Bahasa Indonesia"
        for col, header in enumerate(WORKBOOK_COLUMNS, start=1):
            value = ws.cell(row=FIRST_ROW, column=col).value
            if header == "Confidence":
                assert value == pytest.approx(float(rows[0][header]))
            else:
                assert (value if value is not None else "") == rows[0][header]

    def test_indicator_ids_are_text_cells(self, tmp_path):
        """4.01 and 12.4.1 are different indicators from 4.1 and 12.4: entered
        as numbers they collapse. The template formats E as text; the writer
        must also write a STRING, not a float."""
        rows = [
            wb_row(**{"Indicator ID": "4.01", "Article / Section": "s. 1"}),
            wb_row(**{"Indicator ID": "12.4.1", "Article / Section": "s. 2"}),
            wb_row(**{"Indicator ID": "6.1", "Article / Section": "s. 3"}),
        ]
        _, wb = written(tmp_path, rows)
        ws = wb[SHEET]
        assert [ws.cell(row=r, column=5).value for r in (7, 8, 9)] == ["4.01", "12.4.1", "6.1"]
        for r in (7, 8, 9):
            assert ws.cell(row=r, column=5).data_type == "s"
            assert ws.cell(row=r, column=5).number_format == "@"

    def test_pillar_formula_survives_and_derives_the_pillar(self, tmp_path):
        rows = [wb_row(**{"Indicator ID": "12.4.1", "Article / Section": "s. 2"})]
        _, wb = written(tmp_path, rows)
        ws = wb[SHEET]
        assert ws["O7"].value == (
            '=IF($E7="","",IFERROR(INT($E7),IFERROR(VALUE(LEFT($E7,FIND(".",$E7)-1)),"?")))'
        )
        assert ws["O107"].value.startswith("=IF($E107=")
        assert ws["O108"].value is None
        assert pillar_from_formula(ws["E7"].value) == 12

    def test_example_rows_are_deleted_as_the_instructions_ask(self, tmp_path):
        """Instructions B11: "Delete example rows ... Remove rows 7 and 8
        before submitting." The rows go, and everything below moves up two
        rows the way a spreadsheet application moves it: formulas, row heights,
        validations, the autofilter and the Coverage Matrix bounds."""
        template = load_workbook(template_path())[SHEET]
        result, wb = written(tmp_path, [wb_row()])
        ws = wb[SHEET]
        assert EXAMPLE_ROWS == (7, 8)
        assert result.example_rows_deleted is True
        assert "deleted" in result.note
        snippets = {template.cell(row=r, column=9).value for r in EXAMPLE_ROWS}
        for row in ws.iter_rows():
            for cell in row:
                assert cell.value not in snippets
        # the banner labelled the example rows; they are gone, so its text goes
        # while the row, its merge and its style stay
        assert template["A6"].value.startswith("▸ EXAMPLE ROWS")
        assert ws["A6"].value is None
        assert "A6:O6" in {str(m) for m in ws.merged_cells.ranges}
        assert ws["A6"].font.b == template["A6"].font.b
        assert ws["A6"].fill.fgColor.rgb == template["A6"].fill.fgColor.rgb
        assert ws.row_dimensions[7].height == template.row_dimensions[9].height
        assert ws.row_dimensions[8].height == template.row_dimensions[10].height

    def test_every_formula_reading_output_data_moves_with_the_rows(self, tmp_path):
        template = load_workbook(template_path())
        _, wb = written(tmp_path, [wb_row()])
        def moved(m: re.Match) -> str:
            return re.sub(
                r"\d+", lambda d: str(int(d.group(0)) - 2 if int(d.group(0)) >= 9 else d.group(0)),
                m.group(0),
            )

        def shift(formula: str) -> str:
            return re.sub(r"'Output Data'!\$[A-Z]+\$\d+(?::\$[A-Z]+\$\d+)?", moved, formula)

        n = 0
        for name in template.sheetnames:
            if name == SHEET:
                continue
            for row in template[name].iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and "Output Data" in cell.value:
                        assert wb[name][cell.coordinate].value == shift(cell.value)
                        n += 1
        assert n > 0

    def test_template_structure_is_left_alone(self, tmp_path):
        _, wb = written(tmp_path, [wb_row()])
        assert wb.sheetnames == [
            "Output Data",
            "Indicator Reference",
            "Coverage Matrix",
            "Engine Comparison",
            "Run Record",
            "Submission Checklist",
            "Instructions",
        ]
        ws = wb[SHEET]
        assert ws.auto_filter.ref == "A4:M107"
        assert ws.freeze_panes == "A7"
        validations = {
            str(dv.sqref): (dv.type, dv.formula1) for dv in ws.data_validations.dataValidation
        }
        assert validations["G7:G107"] == ("list", '"NEW,KNOWN"')
        assert validations["N7:N107"] == (
            "list",
            '"English,Thai,Vietnamese,Bahasa Indonesia,Chinese,Hindi,Kazakh,'
            'Russian,Lao,Mongolian,Other"',
        )
        assert validations["L7:L107"][0] == "decimal"
        assert validations["J7:J107"] == ("custom", "lte(LEN(J7),300)")
        # the template's NEW highlight starts one row into the entry area; it
        # moves with it
        assert [str(cf.sqref) for cf in ws.conditional_formatting] == ["G8:G107"]
        assert wb["Coverage Matrix"]["B4"].value == (
            "=COUNTIFS('Output Data'!$O$7:$O$107,1,'Output Data'!$A$7:$A$107,$A4)"
        )
        assert wb["Instructions"]["A1"].value.startswith("How to complete")

    def test_confidence_is_written_as_a_number_and_blanks_stay_blank(self, tmp_path):
        rows = [wb_row(), wb_row(**{"Confidence": "", "Article / Section": "s. 9"})]
        _, wb = written(tmp_path, rows)
        ws = wb[SHEET]
        assert ws["L7"].value == pytest.approx(0.81)
        assert ws["L8"].value is None

    def test_never_writes_past_the_last_entry_row(self, tmp_path):
        rows = [wb_row(**{"Article / Section": f"s. {i}"}) for i in range(150)]
        result, wb = written(tmp_path, rows)
        ws = wb[SHEET]
        assert result.n_rows == ROW_CAP == 101
        assert result.rows_cut == 49
        assert ws.cell(row=LAST_ROW, column=1).value == "Indonesia"
        assert ws.cell(row=LAST_ROW + 1, column=1).value is None


# ---------------------------------------------------------------------------
# Language of Source
# ---------------------------------------------------------------------------


class TestOrganizerLanguage:
    @pytest.mark.parametrize("value", ORGANIZER_LANGUAGES)
    def test_every_organizer_value_maps_to_itself(self, value):
        assert organizer_language(value) == value

    @pytest.mark.parametrize("value", ["Malay", "Bahasa Melayu", "Portuguese", "Tetum", "Khmer"])
    def test_languages_outside_the_list_become_other(self, value):
        assert organizer_language(value) == "Other"

    def test_no_language_at_all_becomes_other(self):
        assert organizer_language(None) == "Other"
        assert organizer_language("") == "Other"

    @pytest.mark.parametrize(
        "raw,expected",
        [
            ("english", "English"),
            ("Indonesian", "Bahasa Indonesia"),
            ("bahasa indonesia", "Bahasa Indonesia"),
            ("Mandarin", "Chinese"),
            ("Lao", "Lao"),
            ("Thai", "Thai"),
        ],
    )
    def test_known_spellings_normalise(self, raw, expected):
        assert organizer_language(raw) == expected


# ---------------------------------------------------------------------------
# the whole export: CSV column, cap, gate, workbook
# ---------------------------------------------------------------------------


HOST = {eco: (p.hosts[0] if p.hosts else f"{eco.lower()}.example.gov") for eco, p in load_portals().items()}


def synth(economy: str, i: int, indicator: str = "12.3", cosine: float = 0.6):
    """One passed MappingRecord on the mappings-only Pillar 12 lane (no absence
    rows), with a unique section so no two rows collide in the battery."""
    doc_id = f"doc_{economy.lower()}_synthetic_act"
    quote = f"A service provider shall register with the authority before offering service {i}."
    rec = MappingRecord(
        mapping_id=f"{doc_id}:c{i:04d}::{indicator}",
        document_id=doc_id,
        chunk_id=f"{doc_id}:c{i:04d}",
        economy=economy,  # type: ignore[arg-type]
        indicator_id=indicator,  # type: ignore[arg-type]
        indicator_name="Licensing scheme for e-commerce providers",
        section=f"s. {i}",
        subsection="(1)",
        verbatim_quote=quote,
        page_number=i + 1,
        impact="Requires registration before offering the service.",
        verification_status="passed",
        controlling_evidence=False,
        extraction_attempts=1,
    )
    return rec, quote, cosine


def synth_dataset(economies, per_economy: int, indicator: str = "12.3"):
    records, texts, cosines, docs = [], {}, {}, {}
    for eco in economies:
        doc_id = f"doc_{eco.lower()}_synthetic_act"
        docs[doc_id] = SyntheticDoc(
            corpus_doc=CorpusDoc(
                economy=eco,
                law_name=f"{eco} Electronic Transactions Act 2021",
                law_number_ref="Act 1 of 2021",
                last_amended="2024",
                source_url=f"https://{HOST[eco]}/act/{eco.lower()}-eta-2021",
            ),
            allow_any_host=True,
        )
        for i in range(per_economy):
            # a spread of cosines wide enough that no two rows of one Economy
            # tie on the two-decimal Confidence, so the rank is unambiguous
            rec, quote, cosine = synth(eco, i, indicator, cosine=0.40 + 0.012 * i)
            records.append(rec)
            texts[rec.chunk_id] = f"PREFIX {quote} SUFFIX"
            cosines[(rec.chunk_id, rec.indicator_id)] = cosine
    return records, texts, cosines, docs


def run_export(tmp_path, records, texts, cosines, docs, **kw):
    return export_all(
        tmp_path,
        records,
        chunk_text_lookup=texts,
        gate_cosine_lookup=cosines,
        coverage_stats={},
        liveness_fn=lambda url: True,
        synthetic_docs=docs,
        **kw,
    )


def csv_rows(path: Path) -> list[dict]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


class TestCsvLanguageColumn:
    def test_language_is_column_14_ahead_of_the_enrichment_columns(self, tmp_path):
        records, texts, cosines, docs = synth_dataset(["SG"], 2)
        result = run_export(tmp_path, records, texts, cosines, docs)
        with result.csv_path.open(encoding="utf-8-sig", newline="") as f:
            header = next(csv.reader(f))
        assert tuple(header[:13]) == COLUMNS
        assert header[13] == LANGUAGE_COLUMN == "Language of Source"
        assert tuple(header[13:]) == EXTRA_COLUMNS

    def test_a_single_language_portal_names_its_language(self, tmp_path):
        records, texts, cosines, docs = synth_dataset(["SG"], 2)
        result = run_export(tmp_path, records, texts, cosines, docs)
        rows = csv_rows(result.csv_path)
        assert {r[LANGUAGE_COLUMN] for r in rows} == {"English"}
        assert all(LANGUAGE_UNKNOWN_NOTE not in r["Notes"] for r in rows)

    def test_a_recorded_document_language_wins_and_maps_to_the_list(self, tmp_path):
        records, texts, cosines, docs = synth_dataset(["MY"], 2)
        meta = {"doc_my_synthetic_act": {"language": "Malay"}}
        result = run_export(tmp_path, records, texts, cosines, docs, document_meta=meta)
        rows = csv_rows(result.csv_path)
        assert {r[LANGUAGE_COLUMN] for r in rows} == {"Other"}
        assert all(LANGUAGE_UNKNOWN_NOTE not in r["Notes"] for r in rows)

    def test_portuguese_and_thai_documents(self, tmp_path):
        records, texts, cosines, docs = synth_dataset(["TH"], 2)
        pt = run_export(
            tmp_path / "pt",
            records,
            texts,
            cosines,
            docs,
            document_meta={"doc_th_synthetic_act": {"language": "Portuguese"}},
        )
        assert {r[LANGUAGE_COLUMN] for r in csv_rows(pt.csv_path)} == {"Other"}
        th = run_export(
            tmp_path / "th",
            records,
            texts,
            cosines,
            docs,
            document_meta={"doc_th_synthetic_act": {"language": "Thai"}},
        )
        assert {r[LANGUAGE_COLUMN] for r in csv_rows(th.csv_path)} == {"Thai"}

    def test_unknown_language_is_other_and_says_so_in_notes(self, tmp_path):
        """Thailand's Portal expects Thai OR English, so an unrecorded
        Language cannot be inferred. The row ships "Other" and discloses it
        rather than guessing at a value criterion C1c is scored on."""
        records, texts, cosines, docs = synth_dataset(["TH"], 2)
        result = run_export(
            tmp_path, records, texts, cosines, docs,
            document_meta={"doc_th_synthetic_act": {"language": None}},
        )
        rows = csv_rows(result.csv_path)
        assert {r[LANGUAGE_COLUMN] for r in rows} == {"Other"}
        assert all(LANGUAGE_UNKNOWN_NOTE in r["Notes"] for r in rows)

    def test_every_shipped_value_is_on_the_organizers_list(self, tmp_path):
        records, texts, cosines, docs = synth_dataset(["SG", "TH", "ID"], 3)
        result = run_export(tmp_path, records, texts, cosines, docs)
        for row in csv_rows(result.csv_path):
            assert row[LANGUAGE_COLUMN] in ORGANIZER_LANGUAGES

    def test_two_exports_of_the_same_records_are_byte_identical(self, tmp_path):
        records, texts, cosines, docs = synth_dataset(["SG", "TH"], 4)
        a = run_export(tmp_path / "a", records, texts, cosines, docs)
        b = run_export(tmp_path / "b", records, texts, cosines, docs)
        assert a.csv_path.read_bytes() == b.csv_path.read_bytes()


class TestRowCap:
    ECONOMIES = ["SG", "MY", "ID", "TH", "LA", "VN"]

    def test_over_cap_ships_101_rows_with_every_economy_present(self, tmp_path):
        records, texts, cosines, docs = synth_dataset(self.ECONOMIES, 25)  # 150
        reviews = {r.mapping_id: "accepted" for r in records}
        result = run_export(tmp_path, records, texts, cosines, docs, reviews=reviews)
        rows = csv_rows(result.csv_path)
        assert len(rows) == ROW_CAP == 101
        assert result.rows_cut == 49
        official = {load_portals()[e].official_name for e in self.ECONOMIES}
        assert {r["Economy"] for r in rows} == official
        # round-robin: no Economy is starved while another takes everything
        per_economy = {e: sum(1 for r in rows if r["Economy"] == e) for e in official}
        assert max(per_economy.values()) - min(per_economy.values()) <= 1

    def test_the_cut_keeps_the_most_confident_rows_of_each_economy(self, tmp_path):
        records, texts, cosines, docs = synth_dataset(self.ECONOMIES, 25)
        reviews = {r.mapping_id: "accepted" for r in records}
        result = run_export(tmp_path, records, texts, cosines, docs, reviews=reviews)
        rows = csv_rows(result.csv_path)
        # synth_dataset raises the cosine with i, so the highest i of an
        # Economy are its most confident rows; the cut must take the tail.
        for economy in {r["Economy"] for r in rows}:
            kept = sorted(
                # "Art." where the Economy drafts in Articles (labels.py)
                int(re.match(r"(?:s|Art)\. (\d+)", r["Article / Section"]).group(1))
                for r in rows
                if r["Economy"] == economy
            )
            assert kept == list(range(25 - len(kept), 25))

    def test_unaccepted_mappings_never_ship_even_under_the_cap(self, tmp_path):
        records, texts, cosines, docs = synth_dataset(self.ECONOMIES, 25)
        reviews = {}
        for i, r in enumerate(records):
            reviews[r.mapping_id] = ("accepted", "rejected", "flagged")[i % 3]
        result = run_export(tmp_path, records, texts, cosines, docs, reviews=reviews)
        rows = csv_rows(result.csv_path)
        accepted = {
            r.verbatim_quote for r in records if reviews[r.mapping_id] == "accepted"
        }
        assert len(rows) == 50 and result.rows_cut == 0
        assert {r["Verbatim Snippet"] for r in rows} == accepted

    def test_select_rows_is_a_no_op_under_the_cap(self):
        rows = [
            {"Economy": "Singapore", "Indicator ID": "6.1", "Article / Section": "s. 1",
             "Confidence": "0.90"},
            {"Economy": "Thailand", "Indicator ID": "6.1", "Article / Section": "s. 2",
             "Confidence": "0.10"},
        ]
        kept, cut = select_rows(rows)
        assert cut == 0 and kept == rows


class TestWorkbookFromExport:
    def test_export_writes_the_workbook_beside_the_csv(self, tmp_path):
        records, texts, cosines, docs = synth_dataset(["SG", "TH"], 3)
        result = run_export(tmp_path, records, texts, cosines, docs)
        assert result.xlsx_path == tmp_path / "submission.xlsx"
        assert result.xlsx_path.exists()

    def test_every_output_data_row_equals_the_csv_row(self, tmp_path):
        records, texts, cosines, docs = synth_dataset(["SG", "TH"], 3)
        result = run_export(tmp_path, records, texts, cosines, docs)
        rows = csv_rows(result.csv_path)
        ws = load_workbook(result.xlsx_path)[SHEET]
        for offset, row in enumerate(rows):
            for col, header in enumerate(WORKBOOK_COLUMNS, start=1):
                cell = ws.cell(row=FIRST_ROW + offset, column=col).value
                expected = row[header]
                if header == "Confidence" and expected:
                    assert cell == pytest.approx(float(expected))
                else:
                    assert (cell if cell is not None else "") == expected
        assert ws.cell(row=FIRST_ROW + len(rows), column=1).value is None

    def test_the_workbook_ships_the_same_rows_as_the_csv_under_the_cap(self, tmp_path):
        records, texts, cosines, docs = synth_dataset(TestRowCap.ECONOMIES, 25)
        reviews = {r.mapping_id: "accepted" for r in records}
        result = run_export(tmp_path, records, texts, cosines, docs, reviews=reviews)
        rows = csv_rows(result.csv_path)
        ws = load_workbook(result.xlsx_path)[SHEET]
        written_ids = [
            ws.cell(row=r, column=5).value for r in range(FIRST_ROW, FIRST_ROW + len(rows))
        ]
        assert len(written_ids) == ROW_CAP
        assert written_ids == [r["Indicator ID"] for r in rows]
        assert ws.cell(row=LAST_ROW + 1, column=1).value is None

    def test_absence_rows_are_csv_only(self, tmp_path):
        """A "No provision found" row is our earned zero and belongs in the
        CSV. In Output Data it would spend one of the 101 entry slots and, via
        the Pillar formula, be counted by the organizers' Coverage Matrix as a
        provision we cited, so the workbook lists provisions only."""
        records, texts, cosines, docs = synth_dataset(["SG"], 2, indicator="6.1")
        result = run_export(tmp_path, records, texts, cosines, docs)
        rows = csv_rows(result.csv_path)
        absence = [r for r in rows if r["Article / Section"] == "No provision found"]
        assert len(absence) == 3  # the Run covered Pillar 6: 6.2, 6.3, 6.4 earned zeros
        ws = load_workbook(result.xlsx_path)[SHEET]
        written = []
        for offset in range(len(rows)):
            value = ws.cell(row=FIRST_ROW + offset, column=6).value
            if value is None:
                break
            written.append(value)
        assert "No provision found" not in written
        assert written == [
            r["Article / Section"] for r in rows if r["Article / Section"] != "No provision found"
        ]

    def test_the_cap_is_budgeted_over_provisions_not_absence_rows(self, tmp_path):
        """Absence rows never reach Output Data, so they must not spend its
        101 slots: counting them would leave entry rows empty while
        battery-green provisions were cut."""
        records, texts, cosines, docs = synth_dataset(["SG"], 120, indicator="6.1")
        result = run_export(tmp_path, records, texts, cosines, docs)
        rows = csv_rows(result.csv_path)
        absence = [r for r in rows if r["Article / Section"] == "No provision found"]
        assert len(absence) == 3  # 6.2, 6.3 and 6.4 earned their zeros
        assert result.rows_cut == 19  # 120 provisions - 101 slots
        assert len(rows) == ROW_CAP + len(absence)
        ws = load_workbook(result.xlsx_path)[SHEET]
        assert ws.cell(row=LAST_ROW, column=1).value == "Singapore"  # the sheet is full
        assert ws.cell(row=LAST_ROW + 1, column=1).value is None
        written = [ws.cell(row=r, column=6).value for r in range(FIRST_ROW, LAST_ROW + 1)]
        assert "No provision found" not in written
        assert written == [
            r["Article / Section"] for r in rows if r["Article / Section"] != "No provision found"
        ]

    def test_the_workbook_can_be_switched_off_for_the_golden_lane(self, tmp_path):
        records, texts, cosines, docs = synth_dataset(["SG"], 2)
        result = run_export(tmp_path, records, texts, cosines, docs, write_xlsx=False)
        assert result.xlsx_path is None
        assert not (tmp_path / "submission.xlsx").exists()


class TestFromTheDatabase:
    """The judge path: a real working database, through export_from_db and the
    CLI, writing both files into the chosen output directory."""

    def _seeded(self, tmp_path):
        from test_pipeline import seed_db_from_goldens

        return seed_db_from_goldens(tmp_path / "db.sqlite")

    def test_export_from_db_writes_both_files(self, tmp_path):
        from regcompass.pipeline import export_from_db

        storage = self._seeded(tmp_path)
        result = export_from_db(storage, tmp_path / "out")
        assert result.csv_path.exists()
        assert result.xlsx_path == tmp_path / "out" / "submission.xlsx"
        assert result.xlsx_path.exists()
        ws = load_workbook(result.xlsx_path)[SHEET]
        assert ws.cell(row=FIRST_ROW, column=1).value == "Singapore"

    def test_a_documents_recorded_language_reaches_the_column(self, tmp_path):
        """The Document's own Language beats the Portal default: Singapore's
        Portal expects English, but a document recorded as Thai ships Thai."""
        from regcompass.pipeline import export_from_db

        storage = self._seeded(tmp_path)
        storage.conn.execute("UPDATE documents SET language = 'Thai'")
        storage.conn.commit()
        assert all(m["language"] == "Thai" for m in storage.document_meta().values())
        result = export_from_db(storage, tmp_path / "out")
        # provision rows: the absence rows name the Portal's default Language
        rows = [
            r for r in csv_rows(result.csv_path)
            if r["Article / Section"] != "No provision found"
        ]
        assert rows and {r[LANGUAGE_COLUMN] for r in rows} == {"Thai"}
        ws = load_workbook(result.xlsx_path)[SHEET]
        assert ws["N9"].value == "Thai"

    def test_cli_export_writes_both_files_to_the_chosen_directory(self, tmp_path, monkeypatch):
        from typer.testing import CliRunner

        from regcompass.cli import app

        monkeypatch.chdir(ROOT)  # config/ resolution
        db = tmp_path / "db.sqlite"
        self._seeded(tmp_path)
        out = tmp_path / "out"
        result = CliRunner().invoke(app, ["export", "--db", str(db), "--out", str(out)])
        assert result.exit_code == 0, result.output
        assert (out / "submission.csv").exists()
        assert (out / "submission.xlsx").exists()
        assert "submission.xlsx" in result.output


# ---------------------------------------------------------------------------
# the Run Record sheet: the steward's check that the hour happened
# ---------------------------------------------------------------------------


def _countif(ws, value: str) -> int:
    """What the organizers' E58/E59 COUNTIF over C12:C56 evaluates to."""
    return sum(1 for r in range(12, 57) if ws.cell(row=r, column=3).value == value)


class TestRunRecordSheet:
    def _sheet(self, n_docs=2, uploads=()):
        from regcompass.workbook import FetchedDocument, RunRecordPass, RunRecordSheet

        return RunRecordSheet(
            pass_a=RunRecordPass("openrouter / openai/gpt-5.6-luna", "10:05", "10:40", 35.0, 1.25),
            pass_b=RunRecordPass("openrouter / qwen/qwen3-30b-a3b-instruct-2507", "10:42", "10:58", 16.0, 0.11),
            documents=[
                FetchedDocument(
                    source_url=f"https://www.pdp.gov.my/act-{i}.pdf",
                    fetched_during="Engine A pass", time_hhmm=f"10:0{i}",
                    size_kb=812 + i, file_type="PDF", title=f"Act {i}",
                )
                for i in range(n_docs)
            ],
            uploads=list(uploads),
        )

    def test_pass_rows_document_log_and_counters(self, tmp_path):
        from regcompass.workbook import RUN_RECORD_SHEET

        result = write_workbook(
            template_path(), tmp_path / "x.xlsx", [], run_record=self._sheet()
        )
        ws = load_workbook(result.path)[RUN_RECORD_SHEET]
        assert [ws.cell(row=5, column=c).value for c in range(1, 7)] == [
            "Engine A — first pass", "openrouter / openai/gpt-5.6-luna",
            "10:05", "10:40", 35, 1.25,
        ]
        assert [ws.cell(row=6, column=c).value for c in range(2, 7)] == [
            "openrouter / qwen/qwen3-30b-a3b-instruct-2507", "10:42", "10:58", 16, 0.11,
        ]
        assert [ws.cell(row=12, column=c).value for c in range(1, 7)] == [
            1, "https://www.pdp.gov.my/act-0.pdf", "Engine A pass", "10:00", 812, "PDF",
        ]
        assert ws["A13"].value == 2 and ws["A14"].value is None
        # the example row is emptied, the organizers' formulas are untouched
        assert all(ws.cell(row=11, column=c).value is None for c in range(1, 7))
        assert ws["E58"].value == '=COUNTIF(C12:C56,"Engine A pass")'
        assert ws["E59"].value == '=COUNTIF(C12:C56,"Engine B pass")'
        assert ws["F7"].value == "=SUM(F5:F6)"
        assert _countif(ws, "Engine A pass") == 2
        assert _countif(ws, "Engine B pass") == 0
        # the short note is the team's own words
        assert all(ws.cell(row=r, column=2).value is None for r in range(62, 68))

    def test_uploads_are_named_but_never_logged_as_fetched(self, tmp_path):
        from regcompass.workbook import RUN_RECORD_SHEET

        result = write_workbook(
            template_path(), tmp_path / "x.xlsx", [],
            run_record=self._sheet(n_docs=1, uploads=["Scanned Act"]),
        )
        ws = load_workbook(result.path)[RUN_RECORD_SHEET]
        assert _countif(ws, "Engine A pass") == 1
        assert "Scanned Act" in ws["A57"].value
        assert "not downloaded" in ws["A57"].value

    def test_no_run_record_leaves_the_sheet_as_the_template_has_it(self, tmp_path):
        from regcompass.workbook import RUN_RECORD_SHEET

        result = write_workbook(template_path(), tmp_path / "x.xlsx", [])
        ws = load_workbook(result.path)[RUN_RECORD_SHEET]
        assert ws["C11"].value == "Engine A pass"  # the organizers' example row
        assert ws["B5"].value is None


class TestRunRecordFromTheDatabase:
    """export_from_db fills the Run Record sheet from the Run Records: the
    Engine A pass is its Discovery plus its Run, the Engine B pass is its Run
    alone, the log lists what the Discovery downloaded, and an upload during
    the hour is not a download."""

    def _hour(self, tmp_path):
        from test_pipeline import seed_db_from_goldens

        storage = seed_db_from_goldens(tmp_path / "db.sqlite")

        def record(run_id, kind, engine, start, end, **finish):
            storage.run_start(
                run_id=run_id, kind=kind, economy="SG",
                pillars=[] if kind == "discovery" else [6, 7],
                indicators=None, engine=engine, started_at=start,
            )
            storage.run_finish(run_id, status="completed", ended_at=end, **finish)

        # 10:00-10:04 Bangkok: Engine A's Discovery downloads two Documents
        record("disc_a", "discovery", None, "2026-10-15T03:00:00.000000Z",
               "2026-10-15T03:04:00.000000Z", documents_fetched=2)
        for i, (url, ct, size) in enumerate([
            ("https://sso.agc.gov.sg/Act/PDPA2012?format=pdf", "application/pdf", 812 * 1024),
            ("https://sso.agc.gov.sg/Act/CA2018", "text/html; charset=utf-8", 300 * 1024),
        ]):
            storage.manifest_add_pending(url, "SG")
            storage.manifest_mark_fetched(
                url, http_status=200, method="httpx", sha256=f"{i}" * 64,
                content_type=ct, size_bytes=size, local_path=f"SG/raw/doc{i}",
            )
            storage.conn.execute(
                "UPDATE crawl_manifest SET fetched_at = ? WHERE url = ?",
                (f"2026-10-15T03:0{i + 1}:30+00:00", url),
            )
        # an index page and a byte-duplicate are not Documents downloaded
        storage.manifest_add_pending("https://sso.agc.gov.sg/Browse", "SG", kind="index")
        storage.manifest_mark_fetched(
            "https://sso.agc.gov.sg/Browse", http_status=200, method="httpx",
            sha256="9" * 64, content_type="text/html", size_bytes=10, local_path="SG/raw/i",
        )
        storage.conn.execute(
            "UPDATE crawl_manifest SET fetched_at = '2026-10-15T03:02:00+00:00'"
            " WHERE url = 'https://sso.agc.gov.sg/Browse'"
        )
        # a file added by hand while the Discovery was still running: its
        # manifest row sits inside the Discovery's window, and it is still
        # not a download
        storage.manifest_add_pending("upload:" + "8" * 64, "SG")
        storage.manifest_mark_fetched(
            "upload:" + "8" * 64, http_status=200, method="manual", sha256="8" * 64,
            content_type="application/pdf", size_bytes=2048, local_path="SG/raw/up",
        )
        storage.conn.execute(
            "UPDATE crawl_manifest SET fetched_at = '2026-10-15T03:03:00+00:00'"
            " WHERE url = ?", ("upload:" + "8" * 64,),
        )
        storage.upsert_document(
            "doc_x", "SG", "8" * 64, local_path="SG/raw/up", source_kind="manual",
        )
        storage.conn.commit()
        # 10:04: the add's own record
        record("disc_upload", "discovery", None, "2026-10-15T03:04:10.000000Z",
               "2026-10-15T03:04:20.000000Z",
               details={"manual": True, "source": "upload", "document_id": "doc_x",
                        "source_url": "https://sso.agc.gov.sg/Act/Scanned"})
        record("run_a", "run", "engine-a", "2026-10-15T03:05:00.000000Z",
               "2026-10-15T03:40:00.000000Z", cost_usd=1.25)
        # the Engine B Run that owns the seeded Mappings
        record("run_one", "run", "engine-b", "2026-10-15T03:42:00.000000Z",
               "2026-10-15T03:58:00.000000Z", cost_usd=0.11)
        return storage

    def test_the_exported_hour_fills_the_sheet(self, tmp_path):
        from regcompass.pipeline import export_from_db
        from regcompass.workbook import RUN_RECORD_SHEET

        storage = self._hour(tmp_path)
        result = export_from_db(storage, tmp_path / "out", run_id="run_one")
        ws = load_workbook(result.xlsx_path)[RUN_RECORD_SHEET]
        assert ws["B5"].value == "openrouter / openai/gpt-5.6-luna"
        assert (ws["C5"].value, ws["D5"].value, ws["E5"].value) == ("10:00", "10:40", 40)
        assert ws["F5"].value == 1.25
        assert ws["B6"].value == "openrouter / qwen/qwen3-30b-a3b-instruct-2507"
        assert (ws["C6"].value, ws["D6"].value, ws["E6"].value) == ("10:42", "10:58", 16)
        assert ws["F6"].value == 0.11
        log = [
            [ws.cell(row=r, column=c).value for c in range(2, 7)]
            for r in range(12, 57) if ws.cell(row=r, column=3).value
        ]
        assert log == [
            ["https://sso.agc.gov.sg/Act/PDPA2012?format=pdf", "Engine A pass", "10:01", 812, "PDF"],
            ["https://sso.agc.gov.sg/Act/CA2018", "Engine A pass", "10:02", 300, "HTML"],
        ]
        assert _countif(ws, "Engine A pass") == 2
        assert _countif(ws, "Engine B pass") == 0
        assert "not downloaded" in ws["A57"].value

    def test_exporting_the_engine_a_run_fills_the_same_sheet(self, tmp_path):
        from regcompass.pipeline import run_record_sheet

        storage = self._hour(tmp_path)
        from_b = run_record_sheet(storage, "run_one")
        from_a = run_record_sheet(storage, "run_a")
        assert from_a == from_b
        assert [d.fetched_during for d in from_a.documents] == ["Engine A pass"] * 2
        assert from_a.uploads == ["https://sso.agc.gov.sg/Act/Scanned"]


class TestRunRecordHourShapes:
    """The hour as it really goes: Runs never fetch, so every fetch of the
    hour is logged once, as the Engine A pass, and the Engine B pass logs
    none, however the Runs overlap or an add lands between them."""

    def _db(self, tmp_path, a_run, b_run, fetches):
        """fetches: [(discovery_id, start, end, details, [(url, fetched_at, path)])]"""
        from test_pipeline import seed_db_from_goldens

        storage = seed_db_from_goldens(tmp_path / "db.sqlite")
        for i, (disc, start, end, details, docs) in enumerate(fetches):
            storage.run_start(run_id=disc, kind="discovery", economy="SG", pillars=[],
                              indicators=None, engine=None, started_at=start)
            storage.run_finish(disc, status="completed", ended_at=end,
                               documents_fetched=len(docs), details=details)
            for j, (url, at, path) in enumerate(docs):
                storage.manifest_add_pending(url, "SG")
                storage.manifest_mark_fetched(
                    url, http_status=200, method="manual" if details else "httpx",
                    sha256=f"{i}{j}".ljust(64, "0"), content_type="application/pdf",
                    size_bytes=4096, local_path=path,
                )
                storage.conn.execute(
                    "UPDATE crawl_manifest SET fetched_at = ? WHERE url = ?", (at, url)
                )
        storage.conn.commit()
        for run_id, engine, (start, end) in (("run_a", "engine-a", a_run), ("run_one", "engine-b", b_run)):
            storage.run_start(run_id=run_id, kind="run", economy="SG", pillars=[6, 7],
                              indicators=None, engine=engine, started_at=start)
            storage.run_finish(run_id, status="completed", ended_at=end)
        return storage

    def _log(self, sheet):
        return [(d.source_url, d.fetched_during) for d in sheet.documents]

    def test_an_add_between_the_runs_is_the_engine_a_pass(self, tmp_path):
        from regcompass.pipeline import run_record_sheet

        storage = self._db(
            tmp_path,
            ("2026-10-15T03:10:00Z", "2026-10-15T03:30:00Z"),
            ("2026-10-15T03:40:00Z", "2026-10-15T03:50:00Z"),
            [
                ("disc_1", "2026-10-15T03:00:00Z", "2026-10-15T03:05:00Z", {},
                 [("https://sso.agc.gov.sg/Act/A", "2026-10-15T03:01:00+00:00", "SG/raw/a.pdf")]),
                ("disc_add", "2026-10-15T03:32:00Z", "2026-10-15T03:33:00Z",
                 {"manual": True, "source": "url"},
                 [("https://sso.agc.gov.sg/Act/B", "2026-10-15T03:32:30+00:00", "SG/raw/b.pdf")]),
            ],
        )
        sheet = run_record_sheet(storage, "run_one")
        assert self._log(sheet) == [
            ("https://sso.agc.gov.sg/Act/A", "Engine A pass"),
            ("https://sso.agc.gov.sg/Act/B", "Engine A pass"),
        ]
        assert (sheet.pass_a.start_hhmm, sheet.pass_a.end_hhmm) == ("10:00", "10:33")
        assert sheet.pass_b.start_hhmm == "10:40"

    def test_runs_that_overlap_log_each_document_once(self, tmp_path):
        from regcompass.pipeline import run_record_sheet

        storage = self._db(
            tmp_path,
            ("2026-10-15T03:10:00Z", "2026-10-15T03:30:00Z"),
            ("2026-10-15T03:15:00Z", "2026-10-15T03:35:00Z"),  # B started while A ran
            [
                ("disc_1", "2026-10-15T03:00:00Z", "2026-10-15T03:05:00Z", {},
                 [("https://sso.agc.gov.sg/Act/A", "2026-10-15T03:01:00+00:00", "SG/raw/a.pdf"),
                  ("https://sso.agc.gov.sg/Act/C", "2026-10-15T03:02:00+00:00", "SG/raw/c.pdf")]),
            ],
        )
        for exported in ("run_a", "run_one"):
            sheet = run_record_sheet(storage, exported)
            assert [d for _, d in self._log(sheet)] == ["Engine A pass"] * 2

    def test_the_prepared_shape_logs_no_engine_b_rows(self, tmp_path):
        """Prepared data: a Discovery, then several adds by address between
        and after the Engine A Run, then the Engine B Run. Every fetch is the
        Engine A pass's; E59 stays at zero."""
        from regcompass.pipeline import export_from_db
        from regcompass.workbook import RUN_RECORD_SHEET

        adds = [
            (f"disc_add{k}", f"2026-10-15T03:3{k}:00Z", f"2026-10-15T03:3{k}:30Z",
             {"manual": True, "source": "url"},
             [(f"https://sso.agc.gov.sg/Act/X{k}", f"2026-10-15T03:3{k}:10+00:00",
               f"SG/raw/x{k}.pdf")])
            for k in range(1, 4)
        ]
        storage = self._db(
            tmp_path,
            ("2026-10-15T03:10:00Z", "2026-10-15T03:30:00Z"),
            ("2026-10-15T03:40:00Z", "2026-10-15T03:50:00Z"),
            [("disc_1", "2026-10-15T03:00:00Z", "2026-10-15T03:05:00Z", {},
              [("https://sso.agc.gov.sg/Act/A", "2026-10-15T03:01:00+00:00", "SG/raw/a.pdf")]),
             *adds],
        )
        result = export_from_db(storage, tmp_path / "out", run_id="run_one")
        ws = load_workbook(result.xlsx_path)[RUN_RECORD_SHEET]
        assert _countif(ws, "Engine A pass") == 4
        assert _countif(ws, "Engine B pass") == 0

    def test_file_type_trusts_the_files_own_suffix(self):
        from regcompass.pipeline import _file_type

        assert _file_type("application/pdf", "SG/raw/page.htm") == "HTML"
        assert _file_type("text/html", "SG/raw/act.pdf") == "PDF"
        assert _file_type("application/pdf", "SG/raw/blob") == "PDF"
