"""Deleting rows the way a spreadsheet application does: every reference to a
row below the deleted block moves up with it, references above it stay."""

from openpyxl import Workbook

from regcompass.sheet_rows import delete_rows, shift_formula


def test_shift_formula_moves_only_rows_below_the_block_on_the_target_sheet():
    f = "=COUNTIFS('Output Data'!$O$9:$O$109,1,'Output Data'!$A$9:$A$109,$A4)"
    assert shift_formula(f, "Coverage Matrix", "Output Data", 7, 2) == (
        "=COUNTIFS('Output Data'!$O$7:$O$107,1,'Output Data'!$A$7:$A$107,$A4)"
    )
    # $A4 is the matrix's own row: another sheet, never moved
    assert shift_formula("=SUM(A1:A20)+B5", "Data", "Data", 7, 2) == "=SUM(A1:A18)+B5"
    assert shift_formula("=SUM(Other!A9:A20)", "Data", "Data", 7, 2) == "=SUM(Other!A9:A20)"
    assert shift_formula("=SUM(A:A)", "Data", "Data", 7, 2) == "=SUM(A:A)"
    assert shift_formula("plain text", "Data", "Data", 7, 2) == "plain text"


def test_delete_rows_moves_values_formulas_heights_and_validations():
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    other = wb.create_sheet("Sum")
    ws["A1"] = "title"
    ws.merge_cells("A1:C1")
    ws["A7"], ws["A8"] = "example 1", "example 2"
    for r in range(9, 13):
        ws.cell(r, 3).value = f"=A{r}*2"
        ws.row_dimensions[r].height = 20 + r
    ws.merge_cells("A12:B12")
    dv = DataValidation(type="custom", formula1="LEN(B9)<5")
    dv.add("B9:B12")
    ws.add_data_validation(dv)
    ws.auto_filter.ref = "A4:C12"
    ws.freeze_panes = "A9"
    other["A1"] = "=SUM(Data!C9:C12)"

    delete_rows(wb, "Data", 7, 2)

    assert ws["A7"].value is None and ws["A1"].value == "title"
    assert ws["C7"].value == "=A7*2" and ws["C10"].value == "=A10*2"
    assert ws["C11"].value is None
    assert [ws.row_dimensions[r].height for r in (7, 8, 9, 10)] == [29, 30, 31, 32]
    assert ws.row_dimensions[11].height is None
    assert {str(m) for m in ws.merged_cells.ranges} == {"A1:C1", "A10:B10"}
    moved = ws.data_validations.dataValidation[0]
    assert str(moved.sqref) == "B7:B10" and moved.formula1 == "LEN(B7)<5"
    assert ws.auto_filter.ref == "A4:C10" and ws.freeze_panes == "A7"
    assert other["A1"].value == "=SUM(Data!C7:C10)"
